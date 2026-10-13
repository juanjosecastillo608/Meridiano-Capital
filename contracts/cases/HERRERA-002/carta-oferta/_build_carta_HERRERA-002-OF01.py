"""Genera 02_Carta_Oferta_HERRERA-002-OF01.docx (BORRADOR INCOMPLETO).

Uso: python _build_carta_HERRERA-002-OF01.py <carpeta_salida>
Luego: soffice --headless --convert-to pdf <docx> --outdir <carpeta_salida>

Los marcadores [PENDIENTE: ...] se resaltan en amarillo para que no pasen
desapercibidos; deben resolverse antes de pasar a LISTO PARA FIRMA.
Paleta: knowledge-base/brand/05-sistema-cromatico.md. Tipografía: alternativas
de sistema oficiales (Cambria / Calibri) de knowledge-base/brand/04-tipografia.md,
sin negrita en titulares (D-040).
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

CODIGO = "HERRERA-002-OF01"
PETROLEO = RGBColor(0x14, 0x31, 0x3A)
GRIS = RGBColor(0x6B, 0x61, 0x54)
TIERRA = RGBColor(0x8B, 0x33, 0x23)
LINEA = "DAD2C0"
DORADO = "C9982E"
FILA_ALT = "F7F4EE"
SERIF = "Cambria"
SANS = "Calibri"

PEND = re.compile(r"(\[PENDIENTE:[^\]]*\])")

doc = Document()
sec = doc.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
sec.top_margin = sec.bottom_margin = Cm(2.2)
sec.left_margin = sec.right_margin = Cm(2.4)

base = doc.styles["Normal"]
base.font.name = SANS
base.font.size = Pt(10.5)
base.element.rPr.rFonts.set(qn("w:eastAsia"), SANS)
base.paragraph_format.space_after = Pt(5)
base.paragraph_format.line_spacing = 1.12


def add_runs(p, text, bold=False, color=None, size=None, font=None, italic=False):
    for chunk in PEND.split(text):
        if not chunk:
            continue
        r = p.add_run(chunk)
        r.bold = bold
        r.italic = italic
        if font:
            r.font.name = font
        if size:
            r.font.size = Pt(size)
        if color is not None:
            r.font.color.rgb = color
        if PEND.fullmatch(chunk):
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
            r.bold = False
    return p


def para(text="", bold=False, color=None, size=None, align=None, after=None, font=None, italic=False, keep=False):
    p = doc.add_paragraph()
    add_runs(p, text, bold=bold, color=color, size=size, font=font, italic=italic)
    if align:
        p.alignment = align
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    if keep:
        p.paragraph_format.keep_with_next = True
    return p


def rule(color=DORADO, size=8):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), str(size))
    b.set(qn("w:color"), color)
    bdr.append(b)
    pPr.append(bdr)
    p.paragraph_format.space_after = Pt(6)
    return p


def heading(n, text):
    p = para(f"{n}. {text}", color=PETROLEO, size=12.5, font=SERIF, after=3, keep=True)
    p.paragraph_format.space_before = Pt(8)
    return p


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear")
    s.set(qn("w:color"), "auto")
    s.set(qn("w:fill"), fill)
    tcPr.append(s)


def borders(table):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:color"), LINEA)
        b.append(e)
    tblPr.append(b)


def table(rows, widths, header=True, total_row=False, size=9.5):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(t)
    for i, row in enumerate(rows):
        tr = t.rows[i]._tr
        trPr = tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit")
        trPr.append(cs)
        is_head = header and i == 0
        is_total = total_row and i == len(rows) - 1
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.width = widths[j]
            c.paragraphs[0].paragraph_format.space_after = Pt(1)
            c.paragraphs[0].paragraph_format.keep_with_next = i < len(rows) - 1
            add_runs(
                c.paragraphs[0],
                val,
                bold=is_head or is_total or (not header and j == 0),
                color=RGBColor(0xFF, 0xFF, 0xFF) if is_head else (PETROLEO if (is_total or (not header and j == 0)) else None),
                size=size,
            )
            if is_head:
                shade(c, "14313A")
            elif i % 2 == 0:
                shade(c, FILA_ALT)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


# ---------- Membrete ----------
para("MERIDIANO CAPITAL", color=PETROLEO, size=20, font=SERIF, after=0)
para("Asesoría, desarrollo y gestión de inversiones inmobiliarias · Asunción, Paraguay", color=GRIS, size=9, after=0)
rule()

p = para("BORRADOR INCOMPLETO — documento de trabajo, no apto para firma hasta resolver los puntos resaltados [PENDIENTE: ...]",
         color=TIERRA, size=8.5, italic=True, after=8)

para("Asunción, 13 de octubre de 2026", align=WD_ALIGN_PARAGRAPH.RIGHT, after=0)
para(f"Ref.: {CODIGO}", color=GRIS, size=9, align=WD_ALIGN_PARAGRAPH.RIGHT, after=10)

para("Señora", after=0)
para("Andrea Elizabeth Cáceres de Ruiz [PENDIENTE: confirmar titular/es registral/es vigente/s y, si corresponde, cónyuge o representante con facultades]", bold=True, after=0)
para("Presente", after=10)

p = para("", after=10)
add_runs(p, "Asunto: ", bold=True, color=PETROLEO)
add_runs(p, "Oferta de compra de la casa ubicada en Gral. Carlos Díaz León N° 3126, Barrio Herrera, Asunción")

para("De mi consideración:", after=6)
para("Quien suscribe, Ana Liesel Guggiari Niederberger, con Cédula de Identidad paraguaya N° 3.209.470, "
     "[PENDIENTE: estado civil, si corresponde], formula por medio de la presente una oferta de compra del inmueble "
     "que se describe a continuación, en los términos y condiciones que se detallan en esta carta.")

# ---------- 1. Inmueble ----------
heading(1, "Inmueble objeto de la oferta")
table([
    ["Tipo", "Casa"],
    ["Ubicación", "Gral. Carlos Díaz León N° 3126, entre Gral. Rogelio Benítez y [PENDIENTE: segunda calle]"],
    ["Barrio / ciudad", "Barrio Herrera, Asunción, Capital"],
    ["Cuenta corriente catastral", "14-0504-18-0"],
    ["Finca / matrícula", "[PENDIENTE: número de finca y distrito según título]"],
    ["Superficie de terreno", "402 m² (según registro municipal; a confirmar con título y plano)"],
    ["Superficie edificada", "112 m² (según registro municipal; a confirmar con título y plano)"],
    ["Bienes incluidos / excluidos", "[PENDIENTE: detallar, o indicar que no hay bienes adicionales]"],
], [Cm(4.6), Cm(11.6)], header=False)

# ---------- 2. Precio ----------
heading(2, "Precio")
para("Precio de venta publicado, como referencia: USD 220.000 (doscientos veinte mil dólares estadounidenses).")
p = para("Precio ofertado: ")
add_runs(p, "Gs 1.000.000.000 (mil millones de guaraníes)", bold=True, color=PETROLEO)
add_runs(p, ", precio total por el inmueble.")
para("La oferta se formula y se pagará íntegramente en guaraníes, moneda en la que la oferente cuenta con un préstamo "
     "bancario pre-aprobado. El precio ofertado es una suma fija en guaraníes y no se ajustará por variaciones del tipo de cambio.")

# ---------- 3. Forma de pago ----------
heading(3, "Forma de pago")
table([
    ["Concepto", "Importe", "Moneda", "Fecha o hito", "Condición"],
    ["Fondos propios de la compradora", "[PENDIENTE: importe]", "Gs", "[PENDIENTE: momento del pago]", "[PENDIENTE]"],
    ["Préstamo bancario pre-aprobado ([PENDIENTE: entidad financiera])", "[PENDIENTE: importe]", "Gs",
     "[PENDIENTE: a la firma de la escritura u otro hito]", "Desembolso de la entidad financiera"],
    ["Total", "1.000.000.000", "Gs", "", ""],
], [Cm(4.6), Cm(2.9), Cm(1.5), Cm(3.6), Cm(3.6)], total_row=True, size=9)
para("[PENDIENTE: indicar si la oferta se acompaña de reserva, seña o anticipo; en su caso: importe, naturaleza, imputación al precio, "
     "momento del pago, receptor autorizado y reglas de devolución o retención]", size=9.5)

# ---------- 4. Condiciones ----------
heading(4, "Condiciones de la oferta")
cond = [
    ("Financiación. ", "La compradora cuenta con un préstamo bancario pre-aprobado en guaraníes. [PENDIENTE: definir si la compra queda "
                       "sujeta a la aprobación definitiva del préstamo y a la tasación de la entidad, y el plazo para obtenerla]"),
    ("Gastos de escritura. ", "Los gastos de escrituración estarán a cargo de la compradora."),
    ("Impuestos de la venta. ", "El impuesto a las transacciones y el IVA que correspondan a la venta estarán a cargo de la parte vendedora."),
    ("Documentación, cargas y ocupación. ", "[PENDIENTE: documentación que aportará la vendedora (título, certificados, informe de condiciones "
                                            "de dominio, impuesto inmobiliario al día), situación de gravámenes y deudas, y ocupación actual]"),
    ("Plazos. ", "[PENDIENTE: plazo para firma de boleto o contrato, si lo hubiera; plazo de escritura; fecha y condiciones de entrega de la posesión]"),
]
for i, (t, body) in enumerate(cond, 1):
    p = para("", after=4)
    p.paragraph_format.left_indent = Cm(0.5)
    add_runs(p, f"4.{i} ", color=PETROLEO)
    add_runs(p, t, bold=True)
    add_runs(p, body)

# ---------- 5. Vigencia ----------
heading(5, "Vigencia y respuesta")
para("La presente oferta se mantiene vigente hasta el [PENDIENTE: fecha] a las [PENDIENTE: hora], hora de Asunción "
     "(America/Asuncion). La respuesta deberá comunicarse por escrito a [PENDIENTE: canal y destinatario de la respuesta]. "
     "Vencido ese plazo sin respuesta, la oferta quedará sin efecto.")

# ---------- 6. Intervención ----------
heading(6, "Intervención comercial")
para("Esta oferta es presentada a la parte vendedora por Meridiano Capital, que interviene como inmobiliaria captadora del lado de "
     "la parte vendedora, a través de su agente Juan José Castillo. La oferente fue presentada por Supremo EAS, que interviene como "
     "inmobiliaria colocadora a través de su agente Mario Luis Dávalos Díaz. Ninguna de las inmobiliarias intervinientes está "
     "facultada, por su sola intervención, para aceptar esta oferta en nombre de la parte vendedora, firmar por la oferente ni recibir fondos.")

# ---------- 7. Naturaleza ----------
heading(7, "Naturaleza y alcance de la oferta")
para("[PENDIENTE: definir el efecto de la aceptación — p. ej., propuesta preliminar sujeta a la firma de un boleto de compraventa, "
     "o compromiso sujeto a condiciones. Requiere confirmación de las partes y revisión jurídica]")

para("Sin otro particular, saludo a usted atentamente.", after=26)

# ---------- Firma oferente ----------
for line in ["______________________________", "Ana Liesel Guggiari Niederberger", "C.I. N° 3.209.470", "Oferente"]:
    para(line, after=0, color=PETROLEO if line.startswith("Ana") else None)

# ---------- Recepción ----------
para("Constancia de recepción", color=PETROLEO, size=13, font=SERIF, after=2).paragraph_format.page_break_before = True
rule(LINEA, 6)
para("Recibí copia de esta oferta. Esta constancia acredita únicamente la recepción y no implica aceptación.", size=10)
table([
    ["Fecha y hora de recepción", "____ / ____ / 2026  ·  ____ : ____ hs (hora de Asunción)"],
    ["Firma", ""],
    ["Aclaración y documento", ""],
], [Cm(5.2), Cm(11.0)], header=False)

# ---------- Respuesta vendedora ----------
para("Respuesta de la parte vendedora", color=PETROLEO, size=13, font=SERIF, after=2).paragraph_format.space_before = Pt(10)
rule(LINEA, 6)
para("Marcar una sola opción:", size=10)
for op in ["☐  Acepto íntegramente la oferta en los términos de esta carta.",
           "☐  Rechazo la oferta.",
           "☐  Formulo la siguiente contraoferta (una contraoferta no implica aceptación de los términos originales):"]:
    para(op, after=3)
box = doc.add_table(rows=1, cols=1)
borders(box)
box.rows[0].height = Cm(3.0)
box.cell(0, 0).width = Cm(16.2)
doc.add_paragraph().paragraph_format.space_after = Pt(2)
para("En caso de aceptación, la parte vendedora reconoce la intervención de Meridiano Capital y de Supremo EAS en esta operación y "
     "asume el pago de la comisión de intermediación del cinco coma cinco por ciento (5,5 %) sobre el precio de venta "
     "[PENDIENTE: IVA incluido o adicional], exigible [PENDIENTE: momento de exigibilidad, p. ej. a la firma de la escritura], "
     "así como el impuesto a las transacciones y el IVA que correspondan a la venta.", size=10)
table([
    ["Firma parte vendedora", ""],
    ["Aclaración y documento", ""],
    ["Asentimiento conyugal o representante", "[PENDIENTE: si corresponde]"],
    ["Fecha y hora", "____ / ____ / 2026  ·  ____ : ____ hs (hora de Asunción)"],
], [Cm(5.2), Cm(11.0)], header=False)

# Pie
footer = sec.footer.paragraphs[0]
add_runs(footer, f"Meridiano Capital · {CODIGO} · Borrador incompleto · v1", color=GRIS, size=8)
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

out = Path(sys.argv[1]) / f"02_Carta_Oferta_{CODIGO}.docx"
doc.save(out)
print(out)
