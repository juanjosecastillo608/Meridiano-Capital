"""Genera la respuesta y contraoferta del caso REPARG-001 (Edificio República Argentina).

Salida: contracts/cases/REPARG-001/client-response/
  Respuesta_Contraoferta_Republica_Argentina_Meridiano.docx
  (el PDF se obtiene con: soffice --headless --convert-to pdf <docx>)

Paleta e identidad Meridiano Capital (knowledge-base/brand/05 y 04). Fraunces nunca
en negrita literal (D-040): el peso del titular sale de tamaño y color.
"""
import io
from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AQUI = Path(__file__).resolve().parent
REPO = AQUI.parents[3]
SALIDA = AQUI.parent / "client-response"
LOGO = REPO / "production/generadores/logo_horizontal.png"
NOMBRE = "Respuesta_Contraoferta_Republica_Argentina_Meridiano.docx"

PET, GOLD, GOLD_D, GREY, TEXTO = "14313A", "C9982E", "A87D22", "6B6154", "2A2620"
LINEA, FILA, CREMA = "DAD2C0", "F7F4EE", "F3EDE3"
SERIF, SANS = "Fraunces", "Poppins"

# ---------- Datos del caso (validados por el usuario, ver Control_Contraoferta) ----------
PRECIO, ENTREGA, SALDO = 820_000, 250_000, 570_000
assert ENTREGA + SALDO == PRECIO


def usd(n):
    return "USD " + f"{n:,}".replace(",", ".")


def rgb(h):
    return RGBColor.from_string(h)


def run(p, text, *, font=SANS, size=10, color=TEXTO, bold=False, italic=False):
    r = p.add_run(text)
    r.font.name = font
    r._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    r.font.size = Pt(size)
    r.font.color.rgb = rgb(color)
    r.bold = bold
    r.italic = italic
    return r


def para(doc_or_cell, text="", *, size=9.5, color=TEXTO, bold=False, italic=False,
         align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=6, font=SANS, keep_next=False):
    p = doc_or_cell.add_paragraph()
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.line_spacing = 1.1
    pf.alignment = align
    pf.keep_with_next = keep_next
    if text:
        run(p, text, font=font, size=size, color=color, bold=bold, italic=italic)
    return p


def shade(cell, fill):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def borders(table, color=LINEA, sz=4, inside=True):
    tblPr = table._element.tblPr
    b = OxmlElement("w:tblBorders")
    edges = ["top", "left", "bottom", "right"] + (["insideH", "insideV"] if inside else [])
    for e in edges:
        el = OxmlElement(f"w:{e}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:color"), color)
        b.append(el)
    tblPr.append(b)


def cell_margins(table, top=60, bottom=60, left=110, right=110):
    tblPr = table._element.tblPr
    m = OxmlElement("w:tblCellMar")
    for k, v in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        el = OxmlElement(f"w:{k}")
        el.set(qn("w:w"), str(v))
        el.set(qn("w:type"), "dxa")
        m.append(el)
    tblPr.append(m)


def cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    trPr.append(el)


def widths(table, cms):
    # LibreOffice ignora el ancho de celda si la tabla no trae grilla y layout fijo.
    table.autofit = False
    tblPr = table._element.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(int(sum(cms) * 567)))
    tblW.set(qn("w:type"), "dxa")
    for col, w in zip(table.columns, cms):
        col.width = Cm(w)
    for row in table.rows:
        for c, w in zip(row.cells, cms):
            c.width = Cm(w)


def keep_together(table):
    # Todas las filas menos la última quedan unidas a la siguiente: la tabla no se parte.
    for row in table.rows[:-1]:
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.keep_with_next = True


def cell_text(c, text, *, size=9, color=TEXTO, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, italic=False):
    c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.alignment = align
    run(p, text, size=size, color=color, bold=bold, italic=italic)
    return p


def rule(doc, color=GOLD, sz=6, before=2, after=8):
    p = para(doc, before=before, after=after)
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"), "single")
    bot.set(qn("w:sz"), str(sz))
    bot.set(qn("w:color"), color)
    bdr.append(bot)
    pPr.append(bdr)
    return p


def field(p, instr, **kw):
    r = run(p, "", **kw)
    for tag, attr in (("w:fldChar", "begin"), ("w:instrText", instr), ("w:fldChar", "end")):
        el = OxmlElement(tag)
        if tag == "w:fldChar":
            el.set(qn("w:fldCharType"), attr)
        else:
            el.set(qn("xml:space"), "preserve")
            el.text = attr
        r._element.append(el)


def logo_bytes():
    im = Image.open(LOGO)
    buf = io.BytesIO()
    im.crop(im.getbbox()).save(buf, format="PNG")
    buf.seek(0)
    return buf


def condicion(doc, n, titulo, cuerpo):
    p = para(doc, after=4)
    p.paragraph_format.left_indent = Cm(0.75)
    p.paragraph_format.first_line_indent = Cm(-0.75)
    run(p, f"{n}.\t", size=9.5, color=GOLD_D, bold=True)
    run(p, titulo + " ", size=9.5, color=PET, bold=True)
    run(p, cuerpo, size=9.5)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(0.75))
    return p


def build():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = SANS
    st.font.size = Pt(10)

    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.top_margin, sec.bottom_margin = Cm(1.9), Cm(1.6)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.header_distance, sec.footer_distance = Cm(0.8), Cm(0.7)

    # ---------- Encabezado: logo + marca de borrador ----------
    hdr = sec.header
    ht = hdr.add_table(rows=1, cols=2, width=Cm(17.0))
    widths(ht, [10.0, 7.0])
    lp = ht.cell(0, 0).paragraphs[0]
    lp.add_run().add_picture(logo_bytes(), height=Cm(1.0))
    rp = ht.cell(0, 1).paragraphs[0]
    rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run(rp, "BORRADOR PARA REVISIÓN", size=7.5, color=GOLD_D, bold=True)
    ht.cell(0, 1).vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    hrule = hdr.add_paragraph()
    hrule.paragraph_format.space_after = Pt(0)
    pPr = hrule._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bot = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "4"), ("color", LINEA)):
        bot.set(qn(f"w:{k}"), v)
    bdr.append(bot)
    pPr.append(bdr)

    # ---------- Pie: contacto institucional + paginación ----------
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(fp, "Meridiano Capital · Asunción, Paraguay · +595 982 853 111 · "
            "juancastillo@meridianocapital.net · www.meridianocapital.net", size=7, color=GREY)
    fp2 = sec.footer.add_paragraph()
    fp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(fp2, "Página ", size=7, color=GREY)
    field(fp2, "PAGE", size=7, color=GREY)
    run(fp2, " de ", size=7, color=GREY)
    field(fp2, "NUMPAGES", size=7, color=GREY)

    # ---------- Lugar, fecha y referencia ----------
    p = para(doc, align=WD_ALIGN_PARAGRAPH.RIGHT, after=1)
    run(p, "Asunción, ", size=9.5, color=GREY)
    run(p, "[fecha de emisión — completar]", size=9.5, color=GOLD_D, bold=True)
    p = para(doc, align=WD_ALIGN_PARAGRAPH.RIGHT, after=8)
    run(p, "Ref.: Propuesta formal de compra del 16 de octubre de 2026 por USD 700.000",
        size=8.5, color=GREY, italic=True)

    # ---------- Título ----------
    para(doc, "RESPUESTA A PROPUESTA DE COMPRA Y CONTRAOFERTA", font=SERIF, size=15.5,
         color=PET, align=WD_ALIGN_PARAGRAPH.LEFT, after=1)
    para(doc, "Edificio República Argentina esquina Tobatí — Asunción", font=SERIF, size=11.5,
         color=GOLD_D, align=WD_ALIGN_PARAGRAPH.LEFT, after=0)
    rule(doc, after=6)

    # ---------- Cuadro de partes ----------
    filas = [
        ("Comprador y oferente", "Matías Manzato — C.I. N.º 9489984"),
        ("Parte vendedora", "M.T. BAT Sociedad Anónima"),
        ("Representante identificada", "Sra. Mabel Torrico"),
        ("Inmueble", "Edificio mixto de departamentos y locales comerciales — Avda. República "
                     "Argentina esquina Tobatí, Asunción — Cuenta Corriente Catastral N.º 14-0322-17"),
        ("Inmobiliaria colocadora", "Meridiano Capital"),
    ]
    t = doc.add_table(rows=len(filas), cols=2)
    borders(t)
    cell_margins(t, top=30, bottom=30)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (k, v) in enumerate(filas):
        shade(t.cell(i, 0), CREMA)
        cell_text(t.cell(i, 0), k.upper(), size=7.5, color=PET, bold=True)
        cell_text(t.cell(i, 1), v, size=8.5)
        cant_split(t.rows[i])
    widths(t, [4.2, 12.8])

    # ---------- Destinatario ----------
    p = para(doc, before=8, after=0, align=WD_ALIGN_PARAGRAPH.LEFT)
    run(p, "Señor ", size=9.5)
    run(p, "Matías Manzato", size=9.5, color=PET, bold=True)
    run(p, " — C.I. N.º 9489984", size=9.5)
    para(doc, "Presente", align=WD_ALIGN_PARAGRAPH.LEFT, after=0)
    para(doc, "Señor Matías Manzato:", before=8, after=5, align=WD_ALIGN_PARAGRAPH.LEFT)

    # ---------- Apertura ----------
    para(doc,
         "En relación con su propuesta formal de compra de fecha 16 de octubre de 2026, por "
         "USD 700.000, correspondiente al edificio ubicado en Avda. República Argentina esquina "
         "Tobatí de Asunción, identificado con Cuenta Corriente Catastral N.º 14-0322-17, propiedad "
         "de M.T. BAT Sociedad Anónima, cuya representante identificada es la señora Mabel Torrico, "
         "Meridiano Capital, en su carácter de inmobiliaria colocadora, canaliza la respuesta "
         "comunicada por la parte vendedora.")
    para(doc,
         "En consideración de los presupuestos y las necesidades de remodelación planteadas por la "
         "parte compradora para la regularización de planos y los arreglos básicos del edificio, la "
         "parte vendedora acepta reducir el precio de venta a USD 820.000 (ochocientos veinte mil "
         "dólares estadounidenses), conforme a las condiciones que se detallan a continuación.")
    para(doc,
         "La reducción toma en cuenta los trabajos planteados por usted; no implica que la parte "
         "vendedora asuma la ejecución ni el costo de dichas reformas o regularizaciones. Esta "
         "comunicación no constituye la aceptación de la propuesta de USD 700.000 ni perfecciona "
         "una compraventa: transmite la contraoferta de la parte vendedora para su consideración.",
         after=6)

    # ---------- 1. Precio y forma de pago ----------
    p = para(doc, after=5, keep_next=True)
    run(p, "1.  ", size=9.5, color=GOLD_D, bold=True)
    run(p, "Precio y forma de pago", size=9.5, color=PET, bold=True)

    pagos = doc.add_table(rows=4, cols=3)
    borders(pagos)
    cell_margins(pagos, top=35, bottom=35)
    pagos.alignment = WD_TABLE_ALIGNMENT.CENTER
    cab = ("Concepto", "Oportunidad", "Importe")
    for j, h in enumerate(cab):
        shade(pagos.cell(0, j), PET)
        cell_text(pagos.cell(0, j), h.upper(), size=7.5, color="FFFFFF", bold=True,
                  align=WD_ALIGN_PARAGRAPH.RIGHT if j == 2 else WD_ALIGN_PARAGRAPH.LEFT)
    filas_pago = [
        ("Entrega inicial a cuenta del precio",
         "A la aceptación de la presente contraoferta", ENTREGA,
         "Doscientos cincuenta mil dólares estadounidenses"),
        ("Saldo de precio",
         "Contra la firma de la escritura pública traslativa de dominio", SALDO,
         "Quinientos setenta mil dólares estadounidenses"),
    ]
    for i, (c, o, imp, letras) in enumerate(filas_pago, start=1):
        if i % 2 == 0:
            for j in range(3):
                shade(pagos.cell(i, j), FILA)
        cp = cell_text(pagos.cell(i, 0), c, size=9, bold=True, color=PET)
        run(pagos.cell(i, 0).add_paragraph(), letras, size=7.5, color=GREY, italic=True)
        cell_text(pagos.cell(i, 1), o, size=9)
        cell_text(pagos.cell(i, 2), usd(imp), size=9.5, bold=True, color=PET,
                  align=WD_ALIGN_PARAGRAPH.RIGHT)
    for j in range(3):
        shade(pagos.cell(3, j), CREMA)
    cell_text(pagos.cell(3, 0), "PRECIO TOTAL DE VENTA", size=8, bold=True, color=PET)
    cell_text(pagos.cell(3, 1), "Ochocientos veinte mil dólares estadounidenses", size=8.5,
              color=GREY, italic=True)
    cell_text(pagos.cell(3, 2), usd(PRECIO), size=10.5, bold=True, color=PET,
              align=WD_ALIGN_PARAGRAPH.RIGHT)
    for r in pagos.rows:
        cant_split(r)
    widths(pagos, [6.0, 7.6, 3.4])
    keep_together(pagos)

    para(doc,
         "La entrega inicial se imputa íntegramente al precio. El saldo de USD 570.000 sustituye "
         "al saldo de USD 450.000 previsto en su propuesta. Las condiciones de la entrega inicial "
         "—receptor, medio de pago y tratamiento— deberán acordarse por escrito entre las partes "
         "antes de su desembolso.", before=5, after=6)

    # ---------- Condiciones ----------
    p = para(doc, after=5, keep_next=True)
    run(p, "Condiciones de la contraoferta", size=9.5, color=PET, bold=True)
    condicion(doc, 2, "Comisión inmobiliaria.",
              "La comisión inmobiliaria será del 5,5 % (cinco coma cinco por ciento), a cargo de la "
              "parte vendedora.")
    condicion(doc, 3, "Gastos de escrituración.",
              "Los gastos de escrituración serán soportados en un 50 % por la parte compradora y en "
              "un 50 % por la parte vendedora. Este reparto se refiere únicamente a los gastos de "
              "escrituración y es independiente de los impuestos indicados en el punto siguiente.")
    condicion(doc, 4, "Impuestos de la operación.",
              "El IVA correspondiente a la venta y los demás impuestos que correspondan a la "
              "transferencia estarán a cargo de la parte vendedora. Estos conceptos no se adicionan "
              "al precio de USD 820.000.")
    condicion(doc, 5, "Acompañamiento de transición.",
              "La parte vendedora se compromete a acompañar a la parte compradora durante seis (6) "
              "meses desde la compra, brindando toda la información y el conocimiento necesarios "
              "para el traspaso de la documentación y de la gestión de locación del edificio. Este "
              "acompañamiento tiene carácter informativo y de transición; no constituye garantía de "
              "renta, administración del edificio, cobranza ni responsabilidad por las decisiones "
              "futuras de la parte compradora.")
    condicion(doc, 6, "Verificación documental.",
              "Conforme a su propuesta, la operación continúa sujeta a la verificación de la "
              "documentación de dominio del inmueble y de que este se encuentre libre de gravámenes. "
              "La presente no implica que dicha verificación ya se haya realizado, ni renuncia "
              "alguna de la parte compradora a revisar el inmueble.")
    condicion(doc, 7, "Formalización.",
              "Las condiciones aquí comunicadas deberán instrumentarse por escrito y ser suscritas "
              "por la parte compradora y por la parte vendedora, esta última por medio de quien "
              "acredite facultades suficientes. Meridiano Capital, en su carácter de inmobiliaria "
              "colocadora, canaliza esta respuesta sin ejercer representación legal de ninguna de "
              "las partes ni obligar a M.T. BAT Sociedad Anónima.")

    # ---------- Vigencia (campo visible pendiente) ----------
    v = doc.add_table(rows=1, cols=1)
    borders(v, color=GOLD, sz=6, inside=False)
    cell_margins(v, top=50, bottom=50, left=140, right=140)
    shade(v.cell(0, 0), FILA)
    vp = v.cell(0, 0).paragraphs[0]
    vp.paragraph_format.space_after = Pt(0)
    run(vp, "Vigencia de esta contraoferta: ", size=9.5, color=PET, bold=True)
    run(vp, "[pendiente de confirmación]", size=9.5, color=GOLD_D, bold=True)
    widths(v, [17.0])

    # ---------- Cierre + firma (bloque indivisible) ----------
    cierre = [
        "Quedamos a su disposición para recibir su conformidad y coordinar con ambas partes los "
        "pasos de instrumentación de la operación, junto con los profesionales que intervengan en "
        "su formalización.",
        "Agradecemos la confianza depositada en Meridiano Capital y le saludamos muy atentamente.",
    ]
    para(doc, cierre[0], before=8, after=4, keep_next=True, align=WD_ALIGN_PARAGRAPH.LEFT)
    para(doc, cierre[1], after=22, keep_next=True, align=WD_ALIGN_PARAGRAPH.LEFT)
    p = para(doc, align=WD_ALIGN_PARAGRAPH.LEFT, after=0, keep_next=True)
    p.paragraph_format.left_indent = Cm(9.8)
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    top = OxmlElement("w:top")
    for k, val in (("val", "single"), ("sz", "6"), ("color", PET), ("space", "4")):
        top.set(qn(f"w:{k}"), val)
    bdr.append(top)
    pPr.append(bdr)
    run(p, "JUAN JOSÉ CASTILLO", size=10.5, color=PET, bold=True)
    for i, (txt, col) in enumerate((("Broker Inmobiliario", TEXTO), ("MERIDIANO CAPITAL", PET),
                                     ("Inmobiliaria colocadora", GOLD_D))):
        q = para(doc, txt, size=9, color=col, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, after=0,
                 keep_next=i < 2)
        q.paragraph_format.left_indent = Cm(9.8)

    SALIDA.mkdir(parents=True, exist_ok=True)
    out = SALIDA / NOMBRE
    doc.save(out)
    return out


if __name__ == "__main__":
    print(build())
