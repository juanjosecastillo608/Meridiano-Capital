#!/usr/bin/env python3
"""Genera archivos de prueba controlados (propiedades FICTICIAS) para las pruebas funcionales.

Ningún dato corresponde a una propiedad real. Requiere: Pillow, python-pptx, openpyxl, python-docx, LibreOffice.
Uso: python make_fixtures.py DESTINO
"""
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont

from docx import Document
from docx.shared import Inches
from openpyxl import Workbook
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches as PI, Pt


def font(size):
    for f in ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/Library/Fonts/Arial.ttf"):
        if Path(f).exists():
            return ImageFont.truetype(f, size)
    return ImageFont.load_default()


def photo(path, w, h, label, hue=(70, 110, 140), exif_orientation=None, fmt=None):
    """Imagen sintética tipo foto (degradé + horizonte + rótulo FIXTURE)."""
    im = Image.new("RGB", (w, h))
    d = ImageDraw.Draw(im)
    for y in range(h):
        t = y / h
        d.line([(0, y), (w, y)], fill=(int(hue[0] + 120 * t), int(hue[1] + 80 * t), int(hue[2] - 60 * t)))
    d.rectangle([w * 0.15, h * 0.45, w * 0.85, h * 0.8], fill=(215, 210, 200), outline=(90, 80, 70), width=max(2, w // 300))
    for i in range(5):
        x = w * (0.2 + i * 0.13)
        d.rectangle([x, h * 0.55, x + w * 0.07, h * 0.7], fill=(60, 70, 80))
    d.text((w * 0.05, h * 0.05), f"FIXTURE · {label} · {w}x{h}", fill="white", font=font(max(14, w // 30)))
    kw = {}
    if exif_orientation:
        ex = Image.Exif()
        ex[0x0112] = exif_orientation
        kw["exif"] = ex
    im.save(path, format=fmt, quality=90, **kw) if (fmt or str(path).endswith((".jpg", ".jpeg"))) else im.save(path, **kw)
    return path


def plano(path, w=1800, h=1200):
    im = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(im)
    d.rectangle([150, 200, 1650, 1000], outline="black", width=4)
    for x in range(150, 1651, 250):
        d.line([(x, 190), (x, 210)], fill="black", width=2)
    d.text((820, 150), "24,00", fill="black", font=font(28))
    d.text((1680, 580), "12,50", fill="black", font=font(28))
    d.text((700, 560), "PLANTA FIXTURE", fill="black", font=font(40))
    d.text((1600, 60), "N ↑", fill="black", font=font(40))
    d.text((150, 1080), "Esc 1/100 · FIXTURE (no es una propiedad real)", fill="black", font=font(26))
    im.save(path)
    return path


def to_pdf(src, outdir):
    with tempfile.TemporaryDirectory() as prof:
        subprocess.run(["soffice", f"-env:UserInstallation=file://{prof}", "--headless", "--convert-to", "pdf",
                        "--outdir", str(outdir), str(src)], capture_output=True, timeout=180)
    return Path(outdir) / (Path(src).stem + ".pdf")


def pptx_residencial(path, img_land, img_port, deform=False):
    p = Presentation()
    p.slide_width, p.slide_height = PI(13.333), PI(7.5)
    s = p.slides.add_slide(p.slide_layouts[6])
    tb = s.shapes.add_textbox(PI(0.5), PI(0.5), PI(6), PI(1.5)).text_frame
    tb.text = "Departamento Fixture Las Lomas"
    tb.add_paragraph().text = "Venta · USD 185.000 · 92 m² cubiertos"
    if deform:
        s.shapes.add_picture(str(img_land), PI(7), PI(0.5), PI(6), PI(6.5))  # 1600x1000 forzada a 6x6.5 -> deformada
    else:
        s.shapes.add_picture(str(img_land), PI(7), PI(0.5), PI(6), PI(3.75))
    s.notes_slide.notes_text_frame.text = "Fuente: ficha del propietario (FIXTURE)."
    s2 = p.slides.add_slide(p.slide_layouts[6])
    t2 = s2.shapes.add_textbox(PI(0.5), PI(0.5), PI(6), PI(4)).text_frame
    t2.text = "Características"
    for line in ("2 dormitorios, 2 baños", "Cochera cubierta", "Expensas USD 120 por mes", "Contacto: vendedor@inmobiliaria-x.com"):
        t2.add_paragraph().text = line
    s2.shapes.add_picture(str(img_port), PI(8), PI(0.5), PI(3.6), PI(4.8))
    p.save(path)


def pptx_corporativa(path, img):
    p = Presentation()
    p.slide_width, p.slide_height = PI(13.333), PI(7.5)
    s = p.slides.add_slide(p.slide_layouts[6])
    s.shapes.add_textbox(PI(0.5), PI(0.4), PI(9), PI(1)).text_frame.text = "Oficinas Fixture Torre Norte — Alquiler corporativo"
    rows = [("Piso", "m² rentables", "Canon USD/m²", "Canon mensual"), ("8", "420", "18", "7.560"), ("9", "420", "18", "7.560"), ("Total", "840", "", "15.120")]
    tbl = s.shapes.add_table(4, 4, PI(0.5), PI(1.6), PI(7), PI(2)).table
    for r, row in enumerate(rows):
        for c, v in enumerate(row):
            tbl.cell(r, c).text = v
    cd = CategoryChartData()
    cd.categories = ["Piso 8", "Piso 9"]
    cd.add_series("m² rentables", (420, 420))
    s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, PI(8), PI(1.6), PI(4.8), PI(3), cd)
    s.shapes.add_picture(str(img), PI(0.5), PI(4.0), PI(4.8), PI(3.0))
    p.save(path)


def main():
    out = Path(sys.argv[1] if len(sys.argv) > 1 else "fixtures")
    if out.exists():
        shutil.rmtree(out)
    img = out / "_img"
    img.mkdir(parents=True)
    land = photo(img / "exterior.jpg", 1600, 1000, "exterior")
    port = photo(img / "interior_vertical.jpg", 900, 1200, "interior", hue=(120, 100, 80))
    # 1) presentación residencial existente
    (out / "01_residencial").mkdir()
    pptx_residencial(out / "01_residencial" / "Depto_Las_Lomas.pptx", land, port)
    # 2) presentación corporativa existente (tabla + gráfico)
    (out / "02_corporativa").mkdir()
    pptx_corporativa(out / "02_corporativa" / "Torre_Norte.pptx", land)
    # 3) PDF con fotografías y precios
    d3 = out / "03_pdf_desarrollo"
    d3.mkdir()
    doc = Document()
    doc.add_heading("Desarrollo Fixture Mburucuyá — Preventa", 1)
    doc.add_paragraph("Lista de precios vigente al 01/10/2026. Precios en USD, IVA incluido.")
    t = doc.add_table(rows=3, cols=3)
    for r, row in enumerate([("Tipología", "m²", "Precio USD"), ("1 dormitorio", "45", "72.000"), ("2 dormitorios", "68", "104.000")]):
        for c, v in enumerate(row):
            t.cell(r, c).text = v
    doc.add_paragraph("Entrega estimada: diciembre de 2028.")
    doc.add_picture(str(land), width=Inches(5))
    doc.save(d3 / "Mburucuya.docx")
    to_pdf(d3 / "Mburucuya.docx", d3)
    (d3 / "Mburucuya.docx").unlink()
    # 4) planilla con tipologías y formas de pago (con fórmula)
    d4 = out / "04_planilla"
    d4.mkdir()
    wb = Workbook()
    ws = wb.active
    ws.title = "Tipologias"
    ws.append(["Tipología", "m²", "Precio USD", "Anticipo 30%", "Cuotas (24)", "Cuota USD"])
    ws.append(["1D", 45, 72000, "=C2*0.3", 24, "=(C2-D2)/E2"])
    ws.append(["2D", 68, 104000, "=C3*0.3", 24, "=(C3-D3)/E3"])
    wb.save(d4 / "Mburucuya_precios.xlsx")
    # 5) plano + fotografías (modo C)
    d5 = out / "05_plano_fotos"
    d5.mkdir()
    plano(d5 / "plano_local.png")
    shutil.copy(land, d5 / "fachada.jpg")
    shutil.copy(port, d5 / "interior.jpg")
    (d5 / "datos.csv").write_text("campo;valor\nsuperficie;300 m²\ncanon;USD 3.600 + IVA por mes\n", encoding="utf-8")
    # 6) datos contradictorios: ficha (rank 5) vs lista de precios (rank 2); dos planillas de igual rank con distinta superficie
    d6 = out / "06_contradicciones"
    d6.mkdir()
    doc = Document()
    doc.add_paragraph("Casa Fixture San Bernardino. Precio: USD 185.000. Superficie: 240 m².")
    doc.save(d6 / "ficha.docx")
    for name, sup in (("lista_a.xlsx", 240), ("lista_b.xlsx", 255)):
        wb = Workbook()
        wb.active.append(["Precio USD", "Superficie m²"])
        wb.active.append([179000, sup])
        wb.save(d6 / name)
    # 7) fotografías con distintas orientaciones y formatos
    d7 = out / "07_orientaciones"
    d7.mkdir()
    photo(d7 / "vertical_exif6.jpg", 1200, 900, "rotada EXIF 6", exif_orientation=6)
    photo(d7 / "cuadrada.webp", 1000, 1000, "cuadrada", fmt="WEBP")
    photo(d7 / "panoramica.tif", 2400, 800, "panorámica", fmt="TIFF")
    photo(d7 / "baja_resolucion.jpg", 480, 320, "baja res")
    # 8) presentación con imagen deformada
    (out / "08_deformada").mkdir()
    pptx_residencial(out / "08_deformada" / "Deformada.pptx", land, port, deform=True)
    # 9) archivo sin datos comerciales (solo fotos en PDF)
    d9 = out / "09_sin_datos"
    d9.mkdir()
    doc = Document()
    doc.add_paragraph("Fotos de la propiedad.")
    doc.add_picture(str(land), width=Inches(5))
    doc.save(d9 / "fotos.docx")
    to_pdf(d9 / "fotos.docx", d9)
    (d9 / "fotos.docx").unlink()
    # 10) sin archivos: carpeta vacía
    (out / "10_sin_archivos").mkdir()
    # extra: archivo dañado
    (out / "11_danado").mkdir()
    data = (out / "01_residencial" / "Depto_Las_Lomas.pptx").read_bytes()
    (out / "11_danado" / "roto.pptx").write_bytes(data[: len(data) // 3])
    shutil.rmtree(img)
    print(f"OK: fixtures en {out}")


if __name__ == "__main__":
    main()
