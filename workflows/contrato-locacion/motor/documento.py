"""Render del modelo (Jinja2 + marcado simple) a .docx con el membrete de Meridiano Capital."""
import copy
import re
from pathlib import Path

import jinja2
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH as A, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm

AQUI = Path(__file__).resolve().parent.parent
SHELL = AQUI / 'plantilla' / 'shell_meridiano.docx'
MODELO = AQUI / 'modelo' / 'MODELO_CONTRATO_LOCACION_v1.txt'
TOKEN = re.compile(r'(\*\*.+?\*\*|⟦.+?⟧)')

SECTORES_BASE = ['SALA – LIVING', 'COMEDOR', 'DORMITORIO PRINCIPAL', 'COCINA', 'BAÑO',
                 'LAVANDERÍA / ÁREA DE SERVICIO', 'GENERAL / TECNOLOGÍA']
# Lista de verificación del modelo (sin cantidades ni estados: se completan al relevar).
CHECKLIST_GENERAL = ['Router WiFi (si se incluye)', 'Extintores de incendio', 'Detector de humo',
                     'Interruptor diferencial / Disyuntor', 'Llaves de vivienda', 'Tarjetas, tags o credenciales de acceso',
                     '{COCHERA_CONTROL}', '{COCHERA}', 'Paredes, pintura y cielorrasos', 'Pisos, zócalos y revestimientos',
                     'Puertas, cerraduras, ventanas y vidrios', 'Balcón, drenajes y barandas']
COLS_INV = ['ARTÍCULO DESCRIPCIÓN', 'MARCA', 'MODELO / N° SERIE', 'CANTIDAD', 'ESTADO', 'PRUEBA FUNCIONAL', 'OBSERVACIONES']


def _shade(cell, fill='D9E2F3'):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), fill); tcPr.append(sh)


def _bordes(t):
    tblPr = t._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for lado in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{lado}')
        e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:space'), '0'); e.set(qn('w:color'), '808080')
        b.append(e)
    tblPr.append(b)


def nuevo_documento():
    d = Document(str(SHELL))
    body = d.element.body
    for el in list(body):
        if not el.tag.endswith('}sectPr'):
            body.remove(el)
    return d


def _estilo(d, nombre, alt='Normal'):
    try:
        return d.styles[nombre]
    except KeyError:
        return d.styles[alt]


def _runs(p, texto, size=None, bold=False):
    for parte in TOKEN.split(texto):
        if not parte:
            continue
        if parte.startswith('**') and parte.endswith('**'):
            _runs(p, parte[2:-2], size, True)
            continue
        r = p.add_run(parte)
        r.bold = bold or None
        if parte.startswith('⟦'):
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW; r.bold = True
        if size:
            r.font.size = Pt(size)


def parrafo(d, texto, align=A.JUSTIFY, size=None, bold=False, estilo='Cuerpo', after=6):
    p = d.add_paragraph(style=_estilo(d, estilo))
    p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    _runs(p, texto, size, bold)
    return p


def _tabla(d, cab, filas, anchos=None, size=8.5):
    t = d.add_table(rows=1, cols=len(cab)); _bordes(t)
    for i, h in enumerate(cab):
        c = t.rows[0].cells[i]; c.text = ''
        r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(size); _shade(c)
    for f in filas:
        cs = t.add_row().cells
        for i, v in enumerate(f):
            cs[i].text = ''
            _runs(cs[i].paragraphs[0], str(v), size)
    if anchos:
        for row in t.rows:
            for i, w in enumerate(anchos):
                row.cells[i].width = Cm(w)
    d.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def _firmas(d, c, solo_locatario=False):
    izq = ['______________________________', c['prop_razon_social'], f"Representada por {c['repr_nombre']}",
           f"RUC N.º {c['prop_ruc']} / C.I. N.º {c['repr_ci']}", 'EL PROPIETARIO']
    der = ['______________________________', c['inq_nombre'].upper() if not c['inq_nombre'].startswith('⟦') else c['inq_nombre'],
           c['inq_documento'], 'EL LOCATARIO']
    d.add_paragraph()
    t = d.add_table(rows=1, cols=2)
    bloques = [[], der] if solo_locatario else [izq, der]
    for cell, lineas in zip(t.rows[0].cells, bloques):
        cell.text = ''
        for i, ln in enumerate(lineas):
            p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
            p.alignment = A.CENTER; p.paragraph_format.space_after = Pt(0)
            _runs(p, ln, 10, bold=ln in ('EL PROPIETARIO', 'EL LOCATARIO'))
    d.add_paragraph()


def _inventario(d, c):
    inv = c['inventario']
    if inv:
        por_sector = {}
        for it in inv:
            por_sector.setdefault(it.get('SECTOR', '') or 'GENERAL', []).append(it)
        for sector, items in por_sector.items():
            parrafo(d, f'**SECTOR: {sector.upper()}**', A.LEFT)
            filas = [[it.get('ARTICULO DESCRIPCION', ''), it.get('MARCA', ''), it.get('MODELO N SERIE', ''),
                      it.get('CANTIDAD', ''), it.get('ESTADO', ''), '', it.get('OBSERVACIONES', '')] for it in items]
            _tabla(d, COLS_INV, filas + [[''] * 7], [4.2, 1.6, 2.2, 1.4, 1.4, 1.8, 2.4], 8)
        return
    parrafo(d, '⟦PENDIENTE DOCUMENTAL: inventario sin relevar — completar en el inmueble. No se listan bienes no confirmados.⟧', A.LEFT)
    for sector in SECTORES_BASE:
        parrafo(d, f'**SECTOR: {sector}**', A.LEFT)
        if sector.startswith('GENERAL'):
            items = []
            for x in CHECKLIST_GENERAL:
                if '{COCHERA' in x:
                    if not c['tiene_cochera']:
                        continue
                    x = 'Control remoto / tarjeta de cochera' if x == '{COCHERA_CONTROL}' else 'Cochera' + c['cochera_txt'].replace(' y Cochera', '')
                items.append([x] + [''] * 6)
            filas = items + [[''] * 7] * 2
        else:
            filas = [[''] * 7 for _ in range(5)]
        _tabla(d, COLS_INV, filas, [4.2, 1.6, 2.2, 1.4, 1.4, 1.8, 2.4], 8)


def _anexo_ii(d, c):
    puntos = '[....................................]'
    nis = f"ANDE – NIS N.º {c['inm_nis']}"
    filas = [['Servicio', nis, nis], ['Número visible del medidor', puntos, puntos],
             ['Lectura del medidor', puntos + ' kWh', puntos + ' kWh'], ['Fecha', puntos, puntos], ['Hora', puntos, puntos],
             ['Llaves de vivienda', puntos, puntos]]
    if c['tiene_cochera']:
        filas.append(['Llaves de cochera', puntos, puntos])
    filas += [['Llaves de buzón', puntos, puntos], ['Controles / tarjetas / tags', puntos, puntos],
              ['Otros medios de acceso', puntos, puntos], ['Observaciones', puntos, puntos],
              ['Referencia de fotos de medidor', '[Archivo ..... / enlace .........]', '[Archivo ..... / enlace .........]']]
    _tabla(d, ['DATO', 'ENTREGA INICIAL', 'DEVOLUCION FORMAL'], filas, [4.5, 6, 6], 9)


def _anexo_iii(d, c):
    el = ['Fachada / acceso / puerta de entrada', 'Sala / living', 'Comedor', 'Dormitorio principal', 'Cocina y electrodomésticos',
          'Baño, grifería, ducha y drenajes', 'Lavandería y lava-secarropas', 'Pisos, paredes, techos y terminaciones',
          'Puertas, ventanas, cerraduras y balcones', 'Tablero eléctrico, detector, disyuntor y luminarias']
    if c['tiene_cochera']:
        el.append('Cochera' + c['cochera_txt'].replace(' y Cochera', '') + ' y medios de acceso')
    el += ['Medidor ANDE y lectura inicial', 'Otros defectos, pendientes o elementos relevantes']
    puntos = '[....................................]'
    _tabla(d, ['N°', 'SECTOR - ELEMENTO', 'ARCHIVO / FOTO - VIDEO', 'OBSERVACIONES'],
           [[str(i), e, puntos, puntos] for i, e in enumerate(el, 1)], [1, 6, 5, 5], 9)


def _cuenta(d, c):
    for etq, v in c['cuenta_filas']:
        p = parrafo(d, f'{etq}: **{v}**', A.LEFT, estilo='List Paragraph', after=0)
        p.paragraph_format.left_indent = Cm(1)
    d.add_paragraph().paragraph_format.space_after = Pt(0)


def renderizar(c, p, salida):
    env = jinja2.Environment(undefined=jinja2.StrictUndefined, trim_blocks=False, lstrip_blocks=True, keep_trailing_newline=True)
    texto = env.from_string(MODELO.read_text(encoding='utf-8')).render(c=c, p=p)
    d = nuevo_documento()
    directivas = {'[[FIRMAS]]': lambda: _firmas(d, c), '[[FIRMA_LOCATARIO]]': lambda: _firmas(d, c, True),
                  '[[INVENTARIO]]': lambda: _inventario(d, c), '[[ANEXO_II]]': lambda: _anexo_ii(d, c),
                  '[[ANEXO_III]]': lambda: _anexo_iii(d, c), '[[CUENTA]]': lambda: _cuenta(d, c)}
    salto = False
    for linea in texto.split('\n'):
        s = linea.rstrip()
        if not s:
            continue
        n_antes = len(d.paragraphs)
        if s == '[[SALTO]]':
            salto = True
            continue
        if s in directivas:
            directivas[s]()
        elif s.startswith('### '):
            q = parrafo(d, s[4:], A.LEFT, bold=True, after=4); q.paragraph_format.space_before = Pt(8)
            q.paragraph_format.keep_with_next = True
        elif s.startswith('## '):
            parrafo(d, s[3:], A.CENTER, 11, True, after=4)
        elif s.startswith('# '):
            parrafo(d, s[2:], A.CENTER, 14, True, after=4)
        elif s.startswith('- '):
            q = parrafo(d, '– ' + s[2:], A.LEFT, estilo='List Paragraph', after=2); q.paragraph_format.left_indent = Cm(1)
        else:
            parrafo(d, s)
        if salto and len(d.paragraphs) > n_antes:
            d.paragraphs[n_antes].paragraph_format.page_break_before = True
            salto = False
    cp = d.core_properties
    cp.author = 'Meridiano Capital — WF-04'; cp.last_modified_by = 'Meridiano Capital — WF-04'
    cp.title = 'Contrato de alquiler de departamento amoblado'; cp.comments = ''
    d.save(salida)
    return texto


def texto_docx(path):
    d = Document(str(path))
    partes = [p.text for p in d.paragraphs]
    for t in d.tables:
        for r in t.rows:
            partes.append(' | '.join(c.text for c in r.cells))
    return '\n'.join(partes)
