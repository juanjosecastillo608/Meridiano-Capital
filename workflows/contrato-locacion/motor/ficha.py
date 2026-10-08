"""Ficha Maestra de Locación v2: esquema, lectura (.docx/.json) y escritura (plantilla .docx).

La ficha es la ÚNICA fuente de datos del contrato. El lector es compatible con la ficha v1
(la primera usada, LOC-001): las secciones se detectan por los títulos numerados
("2. PROPIETARIO / LOCADOR") y cada tabla de dos columnas se lee como etiqueta | valor.
"""
import json
import re
import unicodedata
from pathlib import Path

# Nivel: C = crítico (sin él no hay contrato), R = requerido por el modelo (placeholder si falta),
#        O = opcional (se omite si falta).
SECCIONES = [
    ('OPERACION', '1. OPERACIÓN', ['OPERACI']),
    ('PROPIETARIO', '2. PROPIETARIO / LOCADOR', ['PROPIETARIO', 'LOCADOR']),
    ('INQUILINO', '3. INQUILINO / LOCATARIO', ['INQUILINO', 'LOCATARIO']),
    ('CODEUDOR', '4. CODEUDOR / GARANTE', ['CODEUDOR', 'GARANTE']),
    ('INMUEBLE', '5. INMUEBLE', ['INMUEBLE']),
    ('CONDICIONES', '6. CONDICIONES ECONÓMICAS Y PLAZO', ['CONDICIONES', 'PLAZO']),
    ('CUENTA', '7. CUENTA DESIGNADA PARA PAGOS', ['CUENTA']),
    ('CONTACTOS', '8. CONTACTOS Y NOTIFICACIONES', ['CONTACTO', 'NOTIFICACION']),
    ('CONTROL', '9. CONTROL DOCUMENTAL', ['CONTROL DOCUMENTAL']),
    ('INVENTARIO', '10. ANEXO I – INVENTARIO (opcional, si ya fue relevado)', ['INVENTARIO']),
]

# (clave, sección, etiqueta, nivel, ayuda/ejemplo, alias de etiqueta aceptados)
CAMPOS = [
    ('id_expediente', 'OPERACION', 'ID expediente', 'O', 'LOC-002', []),
    ('tipo', 'OPERACION', 'Tipo', 'O', 'Alquiler de departamento amoblado', []),
    ('titulo', 'OPERACION', 'Título contractual', 'O', 'Fijo por el modelo: CONTRATO DE ALQUILER DE DEPARTAMENTO AMOBLADO', []),
    ('lugar_firma', 'OPERACION', 'Lugar de firma', 'C', 'Asunción, República del Paraguay', []),
    ('fecha_firma', 'OPERACION', 'Fecha de firma', 'C', 'dd/mm/aaaa', []),

    ('prop_razon_social', 'PROPIETARIO', 'Razón social', 'C', 'NOMBRE S.A.', []),
    ('prop_ruc', 'PROPIETARIO', 'RUC', 'C', '80000000-0', []),
    ('prop_domicilio', 'PROPIETARIO', 'Domicilio', 'R', 'Calle N°, Barrio, Ciudad, País', []),
    ('repr_nombre', 'PROPIETARIO', 'Representante', 'C', 'Nombre y apellido', []),
    ('repr_tratamiento', 'PROPIETARIO', 'Tratamiento del representante', 'O', 'Sr. / Sra. (opcional)', []),
    ('repr_ci', 'PROPIETARIO', 'C.I.', 'C', '0.000.000', ['CI', 'C.I. representante']),
    ('repr_caracter', 'PROPIETARIO', 'Carácter', 'C', 'Síndico / Presidente / Apoderado', []),
    ('repr_facultades', 'PROPIETARIO', 'Facultades', 'C', 'Acta/Poder, fecha, Escritura N°, Folio, Sección, Registro', []),
    ('prop_email_notif', 'PROPIETARIO', 'E-mail para notificaciones', 'R', 'correo@dominio', []),
    ('prop_tel_notif', 'PROPIETARIO', 'WhatsApp/teléfono para notificaciones', 'R', '+595 ...', []),

    ('inq_nombre', 'INQUILINO', 'Nombre completo', 'C', 'Nombres y apellidos', []),
    ('inq_tratamiento', 'INQUILINO', 'Tratamiento', 'O', 'Sr. / Sra. (opcional; no se deduce del nombre)', []),
    ('inq_nacionalidad', 'INQUILINO', 'Nacionalidad', 'R', 'País o gentilicio', []),
    ('inq_estado_civil', 'INQUILINO', 'Estado civil', 'O', '', []),
    ('inq_fecha_nac', 'INQUILINO', 'Fecha de nacimiento', 'O', 'dd/mm/aaaa', []),
    ('inq_lugar_nac', 'INQUILINO', 'Lugar de nacimiento', 'O', '', []),
    ('inq_documento', 'INQUILINO', 'Documento', 'C', 'Pasaporte N° X  ó  C.I. N.º X', []),
    ('inq_pais_emisor', 'INQUILINO', 'País emisor', 'O', 'Solo si es pasaporte', []),
    ('inq_nro_personal', 'INQUILINO', 'N° personal', 'O', '', []),
    ('inq_exp_venc', 'INQUILINO', 'Expedición / vencimiento', 'O', 'dd/mm/aaaa / dd/mm/aaaa', []),
    ('inq_ruc', 'INQUILINO', 'RUC', 'O', 'Si tiene', []),
    ('inq_domicilio', 'INQUILINO', 'Domicilio', 'O', '', []),
    ('inq_email', 'INQUILINO', 'E-mail', 'R', '', ['Correo electrónico']),
    ('inq_tel', 'INQUILINO', 'Teléfono', 'R', '', ['WhatsApp/teléfono']),

    ('posee_codeudor', 'CODEUDOR', 'Posee codeudor', 'C', 'SI / NO', []),

    ('inm_edificio', 'INMUEBLE', 'Edificio', 'C', 'Edificio ...', []),
    ('inm_unidad', 'INMUEBLE', 'Unidad', 'C', 'Piso N°, Departamento X', []),
    ('inm_direccion', 'INMUEBLE', 'Dirección', 'C', 'Calle N°, Ciudad, País', []),
    ('inm_ctacte', 'INMUEBLE', 'Cta. Cte. Ctral.', 'R', '00-0000-00', []),
    ('inm_nis', 'INMUEBLE', 'NIS', 'R', 'NIS ANDE', []),
    ('inm_condicion', 'INMUEBLE', 'Condición', 'R', '100% amoblado', []),
    ('inm_cochera', 'INMUEBLE', 'Cochera', 'R', 'N° de cochera, o NO', []),
    ('inm_primera_ocupacion', 'INMUEBLE', 'Primera ocupación', 'R', 'SI / NO', []),
    ('inm_destino', 'INMUEBLE', 'Destino', 'R', 'principalmente a vivienda personal', []),
    ('inm_autoriza_domicilio', 'INMUEBLE', 'Autoriza domicilio fiscal/comercial', 'R', 'SI / NO', []),

    ('canon', 'CONDICIONES', 'Canon mensual', 'C', 'USD 950', []),
    ('iva', 'CONDICIONES', 'IVA', 'R', 'Incluido en el canon', []),
    ('iva_tasa', 'CONDICIONES', 'Tasa de IVA', 'R', '5% (residencial) / 10% (comercial o temporal)', []),
    ('expensas', 'CONDICIONES', 'Expensas', 'R', 'Incluidas en el canon', []),
    ('deposito', 'CONDICIONES', 'Depósito en garantía', 'R', 'USD 950', []),
    ('inicio', 'CONDICIONES', 'Inicio', 'C', 'dd/mm/aaaa', []),
    ('fin', 'CONDICIONES', 'Finalización', 'C', 'dd/mm/aaaa', []),
    ('duracion', 'CONDICIONES', 'Duración', 'C', '12 meses / 1 año', []),
    ('reajuste', 'CONDICIONES', 'Reajuste', 'R', 'No aplica durante la vigencia', []),
    ('vencimiento', 'CONDICIONES', 'Vencimiento mensual', 'C', 'Del día 1 al 5 de cada mes  /  El día 7 de cada mes', []),
    ('primer_pago', 'CONDICIONES', 'Primer pago', 'O', 'p. ej. "a la firma del contrato" (si el inicio cae fuera del vencimiento)', []),
    ('serv_energia', 'CONDICIONES', 'Servicios energéticos', 'R', 'A cargo del inquilino', ['Energía eléctrica']),
    ('internet', 'CONDICIONES', 'Internet', 'R', 'A cargo del inquilino', []),

    ('cta_beneficiario', 'CUENTA', 'Beneficiario', 'C', 'Titular exacto de la cuenta', ['Titular de la cuenta']),
    ('cta_doc_titular', 'CUENTA', 'Documento del titular', 'O', 'C.I. / RUC del titular', []),
    ('cta_moneda', 'CUENTA', 'Moneda / plataforma', 'O', 'USD / Wise', []),
    ('cta_tipo', 'CUENTA', 'Tipo de cuenta', 'O', '', []),
    ('cta_banco', 'CUENTA', 'Banco', 'C', '', []),
    ('cta_dir_banco', 'CUENTA', 'Dirección banco', 'O', '', []),
    ('cta_routing', 'CUENTA', 'Routing (Wire/ACH)', 'O', '', ['Routing']),
    ('cta_numero', 'CUENTA', 'Número de cuenta', 'C', '', []),
    ('cta_swift', 'CUENTA', 'SWIFT/BIC', 'O', '', ['SWIFT', 'Código SWIFT']),
    ('cta_iban', 'CUENTA', 'IBAN', 'O', '', []),

    ('contacto_principal', 'CONTACTOS', 'Contacto de urgencia principal', 'R', 'Nombre · teléfono · e-mail', []),
    ('contacto_alternativo', 'CONTACTOS', 'Contacto de urgencia alternativo', 'R', 'Nombre · teléfono · e-mail', []),
    ('contacto_admin', 'CONTACTOS', 'Administración del edificio', 'R', 'Nombre · teléfono · e-mail', []),
]
CAMPO = {c[0]: c for c in CAMPOS}
PLACEHOLDERS = re.compile(r'^(x+|0+|\[?completar\]?|xx-xxxx-\d+|-+|n/?a|pendiente|\.+)$', re.I)


def _norm(s):
    s = unicodedata.normalize('NFKD', s or '').encode('ascii', 'ignore').decode().upper()
    return re.sub(r'[^A-Z0-9]+', ' ', s).strip()


def _seccion_de(texto):
    m = re.match(r'^\s*(\d+)\s*[\.\)]\s*(.+)$', texto or '')
    if not m:
        return None
    t = _norm(m.group(2))
    for key, _, kws in SECCIONES:
        if any(_norm(k) in t for k in kws):
            return key
    return 'OTRA'


def es_placeholder(v):
    v = (v or '').strip()
    if v.startswith('[') and v.endswith(']'):
        return True
    return v == '' or bool(PLACEHOLDERS.match(v)) or 'COMPLETAR' in v.upper()


def _iter_bloques(doc):
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    for el in doc.element.body.iterchildren():
        if el.tag.endswith('}p'):
            yield Paragraph(el, doc)
        elif el.tag.endswith('}tbl'):
            yield Table(el, doc)


def leer_docx(path):
    """Devuelve {'campos': {clave: valor}, 'anexos': [...], 'inventario': [...], 'no_reconocidos': [...]}."""
    from docx import Document
    doc = Document(path)
    idx = {}
    for clave, sec, etq, *_r, alias in CAMPOS:
        for e in [etq] + alias:
            idx[(sec, _norm(e))] = clave
    out = {'campos': {}, 'anexos': [], 'inventario': [], 'no_reconocidos': [], 'estado_declarado': '', 'reglas': []}
    sec = None
    for b in _iter_bloques(doc):
        if hasattr(b, 'text'):
            t = b.text.strip()
            s = _seccion_de(t)
            if s:
                sec = s
            elif t.upper().startswith('ESTADO:'):
                out['estado_declarado'] = t
            elif sec == 'OTRA' and t:
                out['reglas'].extend(x.strip(' •\t') for x in t.split('\n')
                                     if x.strip(' •\t') and not x.strip().startswith('['))
            continue
        filas = [[c.text.strip() for c in r.cells] for r in b.rows]
        if not filas:
            continue
        if sec == 'CONTROL':
            for f in filas[1:]:
                if len(f) >= 2:
                    out['anexos'].append({'documento': f[0], 'estado': f[1], 'observacion': f[2] if len(f) > 2 else ''})
            continue
        if sec == 'INVENTARIO':
            cab = [_norm(x) for x in filas[0]]
            for f in filas[1:]:
                fila = dict(zip(cab, f))
                if any(v.strip() for v in f):
                    out['inventario'].append(fila)
            continue
        for f in filas:
            if len(f) < 2:
                continue
            etq, val = f[0], f[-1]
            clave = idx.get((sec, _norm(etq)))
            if clave:
                out['campos'][clave] = val
            elif etq and _norm(etq) not in ('CAMPO', 'DATO', 'REGLA DE GENERACION', 'REGLA CONTRACTUAL') and sec not in ('OTRA', None):
                out['no_reconocidos'].append(f'{sec}: {etq}')
    return out


def leer(path):
    p = Path(path)
    if p.suffix.lower() == '.json':
        d = json.loads(p.read_text(encoding='utf-8'))
        d.setdefault('anexos', []); d.setdefault('inventario', []); d.setdefault('no_reconocidos', [])
        d.setdefault('estado_declarado', ''); d.setdefault('reglas', [])
        return d
    return leer_docx(p)


def escribir_docx(path, valores=None, titulo_extra=''):
    """Genera la Ficha Maestra v2 (.docx): en blanco (plantilla) o prellenada con `valores`."""
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    valores = valores or {}
    d = Document()
    for s in d.sections:
        s.left_margin = s.right_margin = Cm(2)
    st = d.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(10)

    def shade(cell, fill):
        tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement('w:shd')
        sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), fill); tcPr.append(sh)

    p = d.add_paragraph(); r = p.add_run('FICHA MAESTRA DE LOCACIÓN — MERIDIANO (v2)'); r.bold = True; r.font.size = Pt(14)
    if titulo_extra:
        p = d.add_paragraph(); r = p.add_run(titulo_extra); r.bold = True
    d.add_paragraph('ESTADO: FICHA_EN_CONSTRUCCION').runs[0].bold = True
    p = d.add_paragraph('Completar solo con datos respaldados por documentación o confirmados. Dejar vacío lo que no se sabe: '
                        'nunca escribir un dato supuesto. Niveles: C = crítico (sin él no se genera contrato) · '
                        'R = requerido por el modelo (si falta, el contrato sale con el punto resaltado) · O = opcional.')
    p.runs[0].italic = True; p.runs[0].font.size = Pt(8.5)
    for key, titulo, _ in SECCIONES:
        p = d.add_paragraph(); p.paragraph_format.space_before = Pt(10)
        r = p.add_run(titulo); r.bold = True; r.font.color.rgb = RGBColor(0x16, 0x32, 0x3C)
        if key == 'CONTROL':
            t = d.add_table(rows=1, cols=3); t.style = 'Table Grid'
            for i, h in enumerate(['Documento / dato', 'Estado', 'Observación']):
                t.rows[0].cells[i].text = h; shade(t.rows[0].cells[i], 'D9E2F3')
            filas = valores.get('_anexos') or [
                {'documento': 'Anexo I – Inventario', 'estado': 'PENDIENTE_DOCUMENTAL', 'observacion': ''},
                {'documento': 'Anexo V – Reglamento interno', 'estado': 'PENDIENTE_DOCUMENTAL', 'observacion': ''}]
            for a in filas:
                c = t.add_row().cells; c[0].text = a['documento']; c[1].text = a['estado']; c[2].text = a.get('observacion', '')
            continue
        if key == 'INVENTARIO':
            cab = ['Sector', 'Artículo descripción', 'Marca', 'Modelo / N° serie', 'Cantidad', 'Estado', 'Observaciones']
            t = d.add_table(rows=1, cols=len(cab)); t.style = 'Table Grid'
            for i, h in enumerate(cab):
                t.rows[0].cells[i].text = h; shade(t.rows[0].cells[i], 'D9E2F3')
            for _ in range(3):
                t.add_row()
            continue
        t = d.add_table(rows=0, cols=3); t.style = 'Table Grid'
        for clave, sec, etq, nivel, ayuda, _a in CAMPOS:
            if sec != key:
                continue
            c = t.add_row().cells
            c[0].text = etq; c[0].paragraphs[0].runs[0].bold = True; shade(c[0], 'F2F2F2')
            c[1].text = nivel; c[1].paragraphs[0].runs[0].font.size = Pt(8)
            v = valores.get(clave, '')
            c[2].text = v if v else ''
            if not v and ayuda:
                r = c[2].paragraphs[0].add_run(f'[{ayuda}]'); r.italic = True; r.font.size = Pt(8)
                r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
        for row in t.rows:
            row.cells[0].width = Cm(5.5); row.cells[1].width = Cm(0.9); row.cells[2].width = Cm(10.6)
    p = d.add_paragraph(); p.paragraph_format.space_before = Pt(10)
    r = p.add_run('11. INSTRUCCIONES ESPECÍFICAS PARA EL CONTRATO'); r.bold = True; r.font.color.rgb = RGBColor(0x16, 0x32, 0x3C)
    reglas = valores.get('_reglas') or []
    for x in reglas:
        d.add_paragraph('• ' + x)
    if not reglas:
        q = d.add_paragraph('[Una instrucción por línea, p. ej.: "No incorporar a NOMBRE como parte contractual."]')
        q.runs[0].italic = True; q.runs[0].font.size = Pt(8.5)
    d.save(path)
