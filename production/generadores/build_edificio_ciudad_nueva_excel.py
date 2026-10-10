"""Edificio Ciudad Nueva — análisis de alquileres, rentabilidad, proyección y plusvalía (versión final cliente).

Genera projects/edificio-ciudad-nueva/entregables/Edificio_Ciudad_Nueva_Alquileres_Rentabilidad.xlsx con fórmulas.
Fuentes: detalle de alquileres y planilla del propietario (verificados por el founder el 2026-10-03),
superficies informadas por el founder el 2026-10-10 (terreno 11 x 31 m, 340 m²; 1.109 m² construidos),
tipo de cambio del día (BCP, cierre interbancario 09/10/2026; actualización del 10/10/2026), parámetros del founder (IVA 5%/10%, vacancia 3%,
administración 8%) y del motor de Meridiano (mantenimiento 5%, IVA de venta 1,5%).
Regla del founder (2026-10-03): recalcular siempre al tipo de cambio del día y dejar asentada su fecha.
Recalcular después con la skill xlsx (scripts/recalc.py).
"""
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = Path(__file__).resolve().parents[2] / "projects/edificio-ciudad-nueva/entregables/Edificio_Ciudad_Nueva_Alquileres_Rentabilidad.xlsx"
if len(sys.argv) > 1:
    OUT = Path(sys.argv[1])

PETROLEO, CREMA, FILA, LINEA = "14313A", "F3EDE3", "F7F4EE", "DAD2C0"
F = "Arial"
H_FONT = Font(name=F, bold=True, color=CREMA, size=10)
H_FILL = PatternFill("solid", fgColor=PETROLEO)
INPUT = Font(name=F, color="0000FF", size=10)        # dato recibido (editable)
CALC = Font(name=F, color="000000", size=10)         # fórmula
LINK = Font(name=F, color="008000", size=10)         # vínculo a otra hoja
PEND = Font(name=F, color="8B3323", italic=True, size=10)
BOLD = Font(name=F, bold=True, size=10)
TITLE = Font(name=F, bold=True, size=13, color=PETROLEO)
NOTE = Font(name=F, italic=True, size=9, color="6B6154")
PEND_FILL = PatternFill("solid", fgColor="FBEFE9")
thin = Side(style="thin", color=LINEA)
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
GS = '"₲" #,##0'
USD = '"USD" #,##0'
USD2 = '"USD" #,##0.00'
PCT = "0.00%"

wb = Workbook()


def sheet(name, title, widths):
    ws = wb.create_sheet(name)
    ws["A1"] = title
    ws["A1"].font = TITLE
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.sheet_view.showGridLines = False
    return ws


def header(ws, row, labels):
    for i, t in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=t)
        c.font, c.fill, c.border = H_FONT, H_FILL, BOX
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[row].height = 30


def put(ws, ref, value, font=CALC, fmt=None, fill=None):
    c = ws[ref]
    c.value = value
    c.font = font
    c.border = BOX
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    return c




def legend(ws, row):
    ws.cell(row=row, column=1, value="Leyenda: azul = dato o parámetro editable; negro = fórmula; verde = vínculo a otra hoja; "
            "rojo itálica = supuesto o criterio a tener presente.").font = NOTE


def plist(ws, start, items):
    """Escribe filas Parámetro | Valor | Tipo | Fuente; devuelve {clave: referencia absoluta}."""
    refs = {}
    for i, (k, name, val, fmt, tipo, src) in enumerate(items, start):
        put(ws, f"A{i}", name, BOLD)
        put(ws, f"B{i}", val, INPUT, fmt)
        put(ws, f"C{i}", tipo, PEND if tipo.startswith("Supuesto") or tipo.startswith("Criterio") else CALC)
        put(ws, f"D{i}", src, CALC).alignment = Alignment(wrap_text=True)
        refs[k] = f"Parametros!$B${i}"
    return refs


wb.remove(wb.active)
# ---------------------------------------------------------------- Parametros
P = sheet("Parametros", "Parámetros del análisis — todos editables", [46, 16, 26, 70])
legend(P, 2)
header(P, 4, ["Parámetro", "Valor", "Tipo", "Fuente / criterio"])
PR = plist(P, 5, [
    ("precio", "Precio de venta (USD)", 340000, USD, "Dato confirmado", "Founder (WEB ID 143028006-118)."),
    ("tc", "Tipo de cambio del día (Gs por USD)", 5694, '#,##0', "Dato de mercado",
     "Cierre del mercado interbancario del viernes 09/10/2026 (último día hábil al sábado 10/10/2026): G. 5.694. Fuente: BCP, informado por ABC Color 09/10/2026 (mercado minorista ~G. 5.720). Actualización anterior: G. 5.873 (02/10/2026)."),
    ("tc_fecha", "Fecha del tipo de cambio", "09/10/2026", None, "Dato de mercado", "D-100 (founder 03/10/2026): recalcular siempre al tipo de cambio del día, aunque la planilla traiga otro."),
    ("tc_planilla", "Tipo de cambio de la planilla original (referencia histórica)", 6100, '#,##0', "Dato recibido", "Planilla del propietario (sin fecha). Solo para comparar; no se usa en los resultados."),
    ("iva_dpto", "IVA alquiler residencial (departamentos)", 0.05, "0%", "Dato confirmado", "Founder 03/10/2026 + D-001/D-045 (parametros_mercado.json › fiscal)."),
    ("iva_com", "IVA alquiler comercial (local y cocheras)", 0.10, "0%", "Dato confirmado",
     "Founder 03/10/2026 + D-001. Criterio conservador: 'Cochera' y 'Cochera y dpto' tributan 10% (concepto no residencial) hasta identificar el local comercial."),
    ("vac", "Vacancia (% del ingreso bruto)", 0.03, "0%", "Supuesto del founder", "Zona de alto tránsito y demanda (Mercado 4 / Av. Eusebio Ayala). Coincide con vacancia_pct del motor (3%)."),
    ("adm", "Administración (% del ingreso bruto)", 0.08, "0%", "Supuesto del founder", "Founder 03/10/2026 (el default del motor es 10%; para este caso rige 8%)."),
    ("mant", "Mantenimiento (% del ingreso bruto)", 0.05, "0%", "Criterio Meridiano",
     "mantenimiento_pct = 5% (parametros_mercado.json › supuestos_operativos_default). Cubre reparaciones menores, pintura y recambios entre inquilinos; obras mayores (impermeabilización, fachada) se presupuestan aparte tras la inspección técnica."),
    ("fijos", "Gastos fijos anuales informados (USD)", 1200, USD, "Dato confirmado", "Planilla del propietario, verificada (founder 03/10/2026)."),
    ("ire", "Impuesto a la renta estimado (% del resultado)", 0.10, "0%", "Criterio conservador",
     "IRE 10% (titularidad por S.A., Modelo A). Aplicado sobre el resultado operativo sin deducir depreciación del edificio: estimación conservadora; la base real la define el contador."),
    ("iva_venta", "IVA efectivo en la venta futura (% del precio)", 0.015, "0.0%", "Dato confirmado", "D-082/D-083: 30% base imponible × 5% = 1,5% efectivo."),
    ("g_cons", "Crecimiento anual de alquileres — conservador", 0.03, "0.0%", "Supuesto", "Por debajo de la inflación (IPC 12 meses usado para el ajuste fiscal 2026: 4,1%)."),
    ("g_base", "Crecimiento anual de alquileres — base", 0.05, "0.0%", "Supuesto", "Inflación + ajuste gradual en renovaciones de contrato."),
    ("g_opt", "Crecimiento anual de alquileres — optimista", 0.08, "0.0%", "Supuesto", "Convergencia hacia los valores publicados del barrio (ver Proyeccion)."),
    ("a_cons", "Valorización anual del inmueble (USD) — conservador", 0.02, "0.0%", "Supuesto", "Hipótesis de escenario, no garantizada."),
    ("a_base", "Valorización anual del inmueble (USD) — base", 0.035, "0.0%", "Supuesto", "Hipótesis de escenario, no garantizada."),
    ("a_opt", "Valorización anual del inmueble (USD) — optimista", 0.05, "0.0%", "Supuesto", "Hipótesis de escenario, no garantizada."),
    ("mk1_lo", "Mercado 1 dormitorio — valor bajo publicado (Gs/mes)", 3000000, GS, "Dato de mercado (C)", "InfoCasas, alquiler 1 dormitorio en Ciudad Nueva, consulta 10/10/2026 (avisos, no contratos): Gs 3.000.000 (44 m²) a Gs 3.200.000 (51 m²)."),
    ("mk1_hi", "Mercado 1 dormitorio — valor alto publicado (Gs/mes)", 3200000, GS, "Dato de mercado (C)", "Ídem."),
    ("mk2_lo", "Mercado 2 dormitorios — valor bajo publicado (Gs/mes)", 2100000, GS, "Dato de mercado (C)", "InfoCasas, alquiler 2 dormitorios en Ciudad Nueva, consulta 10/10/2026 (77 avisos; desde Gs 2.100.000 por 100 m²)."),
    ("mk2_hi", "Mercado 2 dormitorios — valor alto publicado (Gs/mes)", 6300000, GS, "Dato de mercado (C)", "Ídem."),
    ("mkl_lo", "Mercado local comercial — desde (Gs/mes)", 2700000, GS, "Dato de mercado (C)", "InfoCasas, salones comerciales en Ciudad Nueva, consulta 03/10/2026 (sin cambios relevados al 10/10/2026)."),
    ("frente", "Frente del terreno (m)", 11, "0.00", "Dato confirmado", "Founder 10/10/2026."),
    ("fondo", "Fondo del terreno (m)", 31, "0.00", "Dato confirmado", "Founder 10/10/2026."),
    ("sup_terreno", "Superficie del terreno informada (m²)", 340, "#,##0", "Dato confirmado", "Founder 10/10/2026: 340 m². Nota: 11 × 31 = 341 m² (diferencia de 1 m², redondeo de medidas; se usa la superficie informada)."),
    ("sup_constr", "Superficie construida total (m²)", 1109, "#,##0", "Dato confirmado", "Founder 10/10/2026."),
    ("costo_m2", "Costo de construcción de referencia — edificio de departamentos, calidad básica (USD/m²)", 650, USD, "Criterio Meridiano",
     "knowledge-base/investment/market-intelligence/construction-costs/06-costos-de-construccion.md (categoría A, founder 2026-08-15). Obra nueva terminada, sin terreno."),
    ("comp_precio", "Comparable: edificio en venta en Ciudad Nueva (USD)", 650000, USD, "Dato de mercado (C)", "InfoCasas Ref #GB49D6: edificio de 5 niveles y 9 unidades en Ciudad Nueva, consulta 10/10/2026 (precio de oferta; superficie no publicada)."),
    ("comp_unid", "Comparable: unidades del edificio publicado", 9, "0", "Dato de mercado (C)", "Ídem."),
])

# ---------------------------------------------------------- Datos_Originales
D = sheet("Datos_Originales", "Datos originales — transcripción literal del detalle y la planilla del propietario", [16, 16, 18, 16, 16])
D["A2"] = "Texto y orden conservados tal como aparecen en las imágenes. Sin cálculos en esta hoja."
D["A2"].font = NOTE
D["A4"] = "Imagen 1 — Detalle de alquileres (columnas B, C, D de la captura)"
D["A4"].font = BOLD
header(D, 5, ["Columna B", "Columna C", "Columna D"])
orig = [
    (None, "Edificio de departamentos", None), (None, "Tercer piso", None),
    ("Inquilino", "2 dorm.", 1500000), ("Inquilino", "2 dorm.", 1500000), ("Inquilino", "1 dorm.", 1200000), ("Inquilino", "1 dorm.", 1100000),
    (None, "Segundo piso", None),
    ("Inquilino", "2 dorm.", 1600000), ("Inquilino", "2 dorm.", 1500000), ("Inquilino", "2 dorm.", 1500000), ("Inquilino", "2 dorm.", 1500000), ("Inquilino", "1 dorm.", 1250000),
    (None, "Primer piso", None),
    ("Inquilino", "1 dorm.", 1300000), ("Inquilino", "3 dorm.", 1600000),
    (None, "Planta Baja", None),
    ("Inquilino", "1 dorm.", 1200000), ("Inquilino", "1 dorm.", 900000), ("Inquilino", "Cochera", 700000), ("Inquilino", "Cochera y dpto", 1200000),
    (None, "TOTAL", 19550000), ("Cambio", 6100, 3205),
]
r = 6
for b, c, dval in orig:
    put(D, f"A{r}", b, INPUT)
    put(D, f"B{r}", c, INPUT if not isinstance(c, (int, float)) else INPUT, '#,##0' if isinstance(c, int) else None)
    put(D, f"C{r}", dval, INPUT, GS if (dval and dval > 10000) else ('"₲" #,##0' if dval else None))
    r += 1
D.cell(row=r, column=1, value="Nota: la última fila de la captura muestra 'Cambio | 6100 | ₲ 3.205'; el '₲' delante de 3.205 es de la fuente (es un valor en USD).").font = NOTE

r += 2
D[f"A{r}"] = "Imagen 2 — Planilla 'RENTABILIDAD EDIFICIO DE APARTAMENTOS'"
D[f"A{r}"].font = BOLD
r += 1
header(D, r, ["Precio USD", "Renta Anual", "USD Mensual", "Gs Mensual", "Cambio / USD Anual"])
plan = [
    (340000, "Bruta (11,3%)", 3205, 19550000, "6100 / 38.459"),
    (None, "Impuesto anual", 2764, None, "2.764"),
    (None, "Gastos anuales", 1200, None, "1.200"),
    (None, "Renta Neta", None, None, "34.495"),
    (None, "Renta Anual Neta", "10,15%", None, None),
    ("340.000 | 100% | 34.495", "34.495 | X | 340.000", None, None, "(regla de tres de la fuente)"),
]
for row in plan:
    r += 1
    for j, v in enumerate(row, 1):
        put(D, f"{get_column_letter(j)}{r}", v, INPUT)
D.cell(row=r + 1, column=1, value="Nota: en la planilla, 'Impuesto anual' y 'Gastos anuales' aparecen en la columna 'USD Mensual' pero sus valores son anuales (se repiten en 'USD Anual').").font = NOTE

# -------------------------------------------------------- Alquileres_Unidad
A = sheet("Alquileres_Unidad", "Alquileres por unidad — datos informados con verificación documental",
          [6, 12, 13, 22, 13, 24, 9, 16, 14, 14, 30, 46])
legend(A, 2)
header(A, 4, ["N°", "Ref.", "Piso", "Texto original", "Tipología", "Concepto", "IVA", "Alquiler mensual (Gs)",
              "USD (TC del día)", "IVA mensual (Gs)", "Estado", "Observación"])
units = [
    ("P3-01", "Tercer piso", "2 dorm.", "2 dormitorios", "Departamento", 1500000, ""),
    ("P3-02", "Tercer piso", "2 dorm.", "2 dormitorios", "Departamento", 1500000, ""),
    ("P3-03", "Tercer piso", "1 dorm.", "1 dormitorio", "Departamento", 1200000, ""),
    ("P3-04", "Tercer piso", "1 dorm.", "1 dormitorio", "Departamento", 1100000, ""),
    ("P2-01", "Segundo piso", "2 dorm.", "2 dormitorios", "Departamento", 1600000, ""),
    ("P2-02", "Segundo piso", "2 dorm.", "2 dormitorios", "Departamento", 1500000, ""),
    ("P2-03", "Segundo piso", "2 dorm.", "2 dormitorios", "Departamento", 1500000, ""),
    ("P2-04", "Segundo piso", "2 dorm.", "2 dormitorios", "Departamento", 1500000, ""),
    ("P2-05", "Segundo piso", "1 dorm.", "1 dormitorio", "Departamento", 1250000, ""),
    ("P1-01", "Primer piso", "1 dorm.", "1 dormitorio", "Departamento", 1300000, ""),
    ("P1-02", "Primer piso", "3 dorm.", "3 dormitorios", "Departamento", 1600000, ""),
    ("PB-01", "Planta baja", "1 dorm.", "1 dormitorio", "Departamento", 1200000, ""),
    ("PB-02", "Planta baja", "1 dorm.", "1 dormitorio", "Departamento", 900000, ""),
    ("PB-03", "Planta baja", "Cochera", "No aplica", "Cochera", 700000, "IVA 10% (concepto no residencial)."),
    ("PB-04", "Planta baja", "Cochera y dpto", "Mixto", "Cochera y departamento", 1200000,
     "IVA 10% por criterio conservador. Pendiente: identificar si esta fila corresponde al local comercial."),
]
first = 5
for i, (ref, piso, txt, tip, conc, gs, obs) in enumerate(units):
    rr = first + i
    put(A, f"A{rr}", i + 1)
    put(A, f"B{rr}", ref)
    put(A, f"C{rr}", piso)
    put(A, f"D{rr}", txt, INPUT)
    put(A, f"E{rr}", tip)
    put(A, f"F{rr}", conc)
    put(A, f"G{rr}", f'=IF(F{rr}="Departamento",{PR["iva_dpto"]},{PR["iva_com"]})', LINK, "0%")
    put(A, f"H{rr}", gs, INPUT, GS)
    put(A, f"I{rr}", f"=H{rr}/{PR['tc']}", LINK, USD2)
    put(A, f"J{rr}", f"=H{rr}*G{rr}", CALC, GS)
    put(A, f"K{rr}", "Informado con verificación documental", CALC)
    put(A, f"L{rr}", obs or None, CALC).alignment = Alignment(wrap_text=True)
last = first + len(units) - 1
tot = last + 1
put(A, f"G{tot}", "TOTAL", BOLD)
put(A, f"H{tot}", f"=SUM(H{first}:H{last})", BOLD, GS)
put(A, f"I{tot}", f"=SUM(I{first}:I{last})", BOLD, USD2)
put(A, f"J{tot}", f"=SUM(J{first}:J{last})", BOLD, GS)
A.cell(row=tot + 2, column=1, value="Estado confirmado por el founder el 03/10/2026: datos informados con verificación de documentación. "
       "La 'Ref.' es un identificador de trabajo según el orden del detalle, no la numeración real de las unidades.").font = NOTE
A.cell(row=tot + 3, column=1, value="Sin datos personales de inquilinos: la fuente solo dice 'Inquilino'.").font = NOTE
A.freeze_panes = "A5"
RNG = lambda col: f"Alquileres_Unidad!${col}${first}:${col}${last}"
GS_TOT = f"Alquileres_Unidad!$H${tot}"
IVA_TOT = f"Alquileres_Unidad!$J${tot}"

# ------------------------------------------------------------- Conciliacion
C = sheet("Conciliacion", "Conciliación — composición comercial vs. detalle de alquileres", [46, 16, 16, 16, 14, 50])
legend(C, 2)
header(C, 4, ["Concepto", "Composición comercial", "Filas en el detalle", "Diferencia", "Resultado", "Observación"])
mix = [("Departamentos 1 dormitorio", 6, '"1 dormitorio"', "E"), ("Departamentos 2 dormitorios", 6, '"2 dormitorios"', "E"),
       ("Departamentos 3 dormitorios", 1, '"3 dormitorios"', "E"), ("Local comercial", 1, '"Local comercial"', "F"),
       ("Cochera", 0, '"Cochera"', "F"), ("Cochera y departamento", 0, '"Cochera y departamento"', "F")]
for i, (lab, m, crit, col) in enumerate(mix, 5):
    put(C, f"A{i}", lab, BOLD)
    put(C, f"B{i}", m, INPUT, "0")
    put(C, f"C{i}", f"=COUNTIFS({RNG(col)},{crit})", CALC, "0")
    put(C, f"D{i}", f"=C{i}-B{i}", CALC, "0;-0;0")
    put(C, f"E{i}", f'=IF(D{i}=0,"Coincide","Revisar")', CALC)
put(C, "F8", "El local no figura con ese nombre en el detalle; 'Cochera y dpto' se trata a IVA 10% (criterio conservador).", CALC).alignment = Alignment(wrap_text=True)
put(C, "A12", "Suma recalculada vs. total indicado (Gs)", BOLD)
put(C, "B12", f"={GS_TOT}", LINK, GS)
put(C, "C12", 19550000, INPUT, GS)
put(C, "D12", "=B12-C12", CALC, GS)
put(C, "E12", '=IF(D12=0,"Suma correcta","DIFERENCIA")', CALC)
put(C, "A14", "Subtotal por piso", TITLE)
header(C, 15, ["Piso", "Filas", "Alquiler mensual (Gs)", "USD (TC del día)", "% del total", ""])
for i, piso in enumerate(["Tercer piso", "Segundo piso", "Primer piso", "Planta baja"], 16):
    put(C, f"A{i}", piso, BOLD)
    put(C, f"B{i}", f"=COUNTIFS({RNG('C')},A{i})", CALC, "0")
    put(C, f"C{i}", f"=SUMIFS({RNG('H')},{RNG('C')},A{i})", CALC, GS)
    put(C, f"D{i}", f"=C{i}/{PR['tc']}", CALC, USD2)
    put(C, f"E{i}", f"=C{i}/$C$20", CALC, "0.0%")
put(C, "A20", "Total", BOLD)
for col, fmt in (("B", "0"), ("C", GS), ("D", USD2), ("E", "0.0%")):
    put(C, f"{col}20", f"=SUM({col}16:{col}19)", BOLD, fmt)

# ---------------------------------------------------------- Ingresos_Gastos
G = sheet("Ingresos_Gastos", "Ingresos y egresos anuales — tipo de cambio del día", [52, 18, 14, 66])
legend(G, 2)
header(G, 4, ["Concepto", "USD por año", "% del bruto", "Base de cálculo"])
lines = [
    ("Ingreso mensual (Gs)", f"={GS_TOT}", GS, None, "Suma del detalle de alquileres."),
    ("Ingreso mensual (USD)", f"=B5/{PR['tc']}", USD2, None, "Gs mensual / tipo de cambio del día."),
    ("Ingreso bruto anual", "=B6*12", USD2, "=B7/$B$7", "Ingreso mensual USD × 12."),
    ("IVA (5% departamentos, 10% local y cocheras)", f"=-{IVA_TOT}*12/{PR['tc']}", USD2, "=-B8/$B$7", "IVA por concepto (hoja Alquileres_Unidad), anualizado."),
    ("Vacancia", f"=-B7*{PR['vac']}", USD2, "=-B9/$B$7", "3% del bruto: zona de alto tránsito y demanda."),
    ("Administración", f"=-B7*{PR['adm']}", USD2, "=-B10/$B$7", "8% del bruto."),
    ("Mantenimiento", f"=-B7*{PR['mant']}", USD2, "=-B11/$B$7", "5% del bruto (criterio Meridiano)."),
    ("Gastos fijos informados", f"=-{PR['fijos']}", USD2, "=-B12/$B$7", "Planilla del propietario, verificada."),
    ("Resultado operativo antes de impuesto a la renta", "=SUM(B7:B12)", USD2, "=B13/$B$7", "Bruto − IVA − vacancia − administración − mantenimiento − gastos fijos."),
    ("Impuesto a la renta estimado (IRE 10%)", f"=-B13*{PR['ire']}", USD2, "=-B14/$B$7", "Estimación conservadora (sin depreciación)."),
    ("Resultado neto anual", "=B13+B14", USD2, "=B15/$B$7", "Ingreso neto para el inversor."),
]
for i, (lab, f, fmt, pc_, base) in enumerate(lines, 5):
    bold = i in (7, 13, 15)
    put(G, f"A{i}", lab, BOLD if bold else CALC)
    put(G, f"B{i}", f, BOLD if bold else (LINK if "!" in f else CALC), fmt)
    if pc_:
        put(G, f"C{i}", pc_, CALC, "0.0%")
    put(G, f"D{i}", base, CALC).alignment = Alignment(wrap_text=True)
G["A5"].font = CALC
put(G, "A17", "Lectura mensual del resultado neto (USD)", BOLD)
put(G, "B17", "=B15/12", CALC, USD2)
put(G, "A18", "IVA efectivo sobre el bruto", BOLD)
put(G, "B18", "=-B8/B7", CALC, "0.00%")
BRUTO, PRE, NETO, IVAEF = "Ingresos_Gastos!$B$7", "Ingresos_Gastos!$B$13", "Ingresos_Gastos!$B$15", "Ingresos_Gastos!$B$18"

# -------------------------------------------------------------- Rentabilidad
R_ = sheet("Rentabilidad", "Rentabilidad sobre el precio — tipo de cambio del día", [56, 16, 66])
legend(R_, 2)
header(R_, 4, ["Indicador", "Valor", "Base de cálculo"])
rl = [
    ("Rentabilidad bruta", f"={BRUTO}/{PR['precio']}", "0.00%", "Ingreso bruto anual / precio."),
    ("Rentabilidad neta antes de impuesto a la renta", f"={PRE}/{PR['precio']}", "0.00%", "Resultado operativo / precio."),
    ("Rentabilidad neta final", f"={NETO}/{PR['precio']}", "0.00%", "Resultado neto (después de IVA, vacancia, administración, mantenimiento, gastos fijos e IRE estimado) / precio."),
    ("Múltiplo precio / ingreso bruto anual (veces)", f"={PR['precio']}/{BRUTO}", "0.00", "Años de ingreso bruto equivalentes al precio."),
    ("Ingreso neto mensual (USD)", f"={NETO}/12", USD2, ""),
]
for i, (lab, f, fmt, base) in enumerate(rl, 5):
    put(R_, f"A{i}", lab, BOLD)
    put(R_, f"B{i}", f, CALC, fmt)
    put(R_, f"C{i}", base, CALC).alignment = Alignment(wrap_text=True)
put(R_, "A11", "Comparación con la planilla original del propietario (referencia)", TITLE)
header(R_, 12, ["Indicador", "Valor", "Observación"])
put(R_, "A13", "Ingreso bruto anual al TC de la planilla (USD)", BOLD)
put(R_, "B13", f"={GS_TOT}*12/{PR['tc_planilla']}", CALC, USD2)
put(R_, "C13", "Con Gs 6.100: USD 38.459. Al TC del día el mismo ingreso en guaraníes equivale a más dólares.", CALC)
put(R_, "A14", "Rentabilidad bruta al TC de la planilla", BOLD)
put(R_, "B14", f"=B13/{PR['precio']}", CALC, "0.00%")
put(R_, "A15", "Rentabilidad 'neta' de la planilla (solo impuesto USD 2.764 y gastos USD 1.200)", BOLD)
put(R_, "B15", f"=(B13-2764-1200)/{PR['precio']}", CALC, "0.00%")
put(R_, "C15", "Reemplazada por el cálculo completo de arriba (IVA por concepto, vacancia, administración, mantenimiento, IRE).", CALC)
R_.cell(row=17, column=1, value="Indicadores sobre el precio de venta, sin gastos de adquisición (escribanía, impuestos de transferencia, honorarios). No constituyen rentabilidad garantizada.").font = NOTE

put(R_, "A19", "Sensibilidad al tipo de cambio (mismo ingreso en guaraníes)", TITLE)
header(R_, 20, ["Tipo de cambio (Gs/USD)", "Rent. bruta", "Rent. neta final"])
for i, (lab, tc_) in enumerate([("Planilla original: 6.100", 6100), ("Actualización 02/10/2026: 5.873", 5873), ("Hoy (09/10/2026): TC del día", PR["tc"]), ("Escenario: 6.300", 6300)], 21):
    put(R_, f"A{i}", lab, BOLD)
    tcref = tc_ if isinstance(tc_, str) else tc_
    put(R_, f"B{i}", f"={GS_TOT}*12/{tcref}/{PR['precio']}", CALC, "0.00%")
    put(R_, f"C{i}", f"=(({GS_TOT}*12/{tcref})*(1-{IVAEF}-{PR['vac']}-{PR['adm']}-{PR['mant']})-{PR['fijos']})*(1-{PR['ire']})/{PR['precio']}", CALC, "0.00%")
R_.cell(row=25, column=1, value="Los alquileres se cobran en guaraníes: si el guaraní se aprecia (TC más bajo), la renta en USD sube; si se deprecia, baja.").font = NOTE

# --------------------------------------------------------- Superficies_Valor
SV = sheet("Superficies_Valor", "Superficies y valor del inmueble", [62, 18, 66])
legend(SV, 2)
header(SV, 4, ["Indicador", "Valor", "Base de cálculo"])
sv = [
    ("Terreno: frente × fondo (m²)", f"={PR['frente']}*{PR['fondo']}", "#,##0", "11 × 31 = 341 m² (control aritmético)."),
    ("Superficie de terreno informada (m²)", f"={PR['sup_terreno']}", "#,##0", "Se usa la superficie informada (340 m²)."),
    ("Superficie construida total (m²)", f"={PR['sup_constr']}", "#,##0", "Informada por el founder."),
    ("Índice de construcción (m² construidos / m² de terreno)", "=B7/B6", "0.00", "Veces la superficie del terreno."),
    ("Precio por m² construido (USD, terreno incluido)", f"={PR['precio']}/B7", USD2, "Precio / superficie construida."),
    ("Precio por m² de terreno (USD, construcción incluida)", f"={PR['precio']}/B6", USD2, "Precio / superficie del terreno."),
    ("Costo de construir hoy la misma superficie (USD, sin terreno)", f"=B7*{PR['costo_m2']}", USD, "1.109 m² × USD 650/m² (edificio de departamentos, calidad básica, obra nueva)."),
    ("Precio / costo de construcción nueva", f"={PR['precio']}/B11", "0%", "El precio equivale a este % del costo de levantar hoy un edificio de igual superficie, sin contar el terreno."),
    ("Unidades que generan renta (13 dptos. + local)", 14, "0", "Composición comercial."),
    ("Precio por unidad (USD)", f"={PR['precio']}/B13", USD, "Precio / 14 unidades."),
    ("Comparable: precio por unidad del edificio publicado en Ciudad Nueva (USD)", f"={PR['comp_precio']}/{PR['comp_unid']}", USD, "USD 650.000 / 9 unidades (precio de oferta)."),
    ("Ingreso bruto anual por m² construido (USD)", f"={BRUTO}/B7", USD2, "Ingreso bruto anual / m² construidos."),
]
for i, (lab, f, fmt, base) in enumerate(sv, 5):
    put(SV, f"A{i}", lab, BOLD)
    put(SV, f"B{i}", f, INPUT if isinstance(f, int) else CALC, fmt)
    put(SV, f"C{i}", base, CALC).alignment = Alignment(wrap_text=True)
SV.cell(row=18, column=1, value="Referencias, no tasación: el edificio es usado (antigüedad no informada) y el costo de reposición no descuenta depreciación. "
        "El comparable es un precio de oferta sin superficie publicada.").font = NOTE

# --------------------------------------------------------------- Proyeccion
PJ = sheet("Proyeccion", "Proyección de alquiler a 5 años — tres escenarios (USD, TC del día constante)", [20, 16, 16, 16, 16, 16, 16])
legend(PJ, 2)
header(PJ, 4, ["Año", "Bruto — conservador", "Neto — conservador", "Bruto — base", "Neto — base", "Bruto — optimista", "Neto — optimista"])
GR = [PR["g_cons"], PR["g_base"], PR["g_opt"]]
net_factor = f"(1-{IVAEF}-{PR['vac']}-{PR['adm']}-{PR['mant']})"
for y in range(1, 6):
    r = 4 + y
    put(PJ, f"A{r}", y, BOLD, "0")
    for k, g in enumerate(GR):
        cb, cn = get_column_letter(2 + 2 * k), get_column_letter(3 + 2 * k)
        put(PJ, f"{cb}{r}", f"={BRUTO}*(1+{g})^(A{r}-1)", CALC, USD)
        put(PJ, f"{cn}{r}", f"=({cb}{r}*{net_factor}-{PR['fijos']})*(1-{PR['ire']})", CALC, USD)
put(PJ, "A10", "Total 5 años", BOLD)
for col in "BCDEFG":
    put(PJ, f"{col}10", f"=SUM({col}5:{col}9)", BOLD, USD)
put(PJ, "A11", "Crecimiento anual", BOLD)
for k, g in enumerate(GR):
    put(PJ, f"{get_column_letter(2 + 2 * k)}11", f"={g}", LINK, "0.0%")
PJ.cell(row=12, column=1, value="Supuestos: crecimiento de alquileres en guaraníes; IVA, vacancia, administración y mantenimiento como % del bruto; gastos fijos constantes; "
        "IRE 10% estimado; tipo de cambio del día constante. Proyección ilustrativa, no garantizada.").font = NOTE

put(PJ, "A14", "Referencia de mercado: alquiler actual vs. avisos publicados en Ciudad Nueva", TITLE)
header(PJ, 15, ["Tipología", "Promedio actual (Gs)", "Mercado bajo (Gs)", "Mercado alto (Gs)", "Brecha vs. bajo", "Unidades", ""])
mk = [("1 dormitorio", '"1 dormitorio"', PR["mk1_lo"], PR["mk1_hi"]), ("2 dormitorios", '"2 dormitorios"', PR["mk2_lo"], PR["mk2_hi"])]
for i, (lab, crit, lo, hi) in enumerate(mk, 16):
    put(PJ, f"A{i}", lab, BOLD)
    put(PJ, f"B{i}", f"=AVERAGEIFS({RNG('H')},{RNG('E')},{crit})", CALC, GS)
    put(PJ, f"C{i}", f"={lo}", LINK, GS)
    put(PJ, f"D{i}", f"={hi}", LINK, GS)
    put(PJ, f"E{i}", f"=C{i}/B{i}-1", CALC, "0%")
    put(PJ, f"F{i}", f"=COUNTIFS({RNG('E')},{crit})", CALC, "0")
put(PJ, "A18", "Local comercial (desde)", BOLD)
put(PJ, "C18", f"={PR['mkl_lo']}", LINK, GS)
PJ.cell(row=19, column=1, value="Avisos publicados (InfoCasas, consulta 10/10/2026, categoría C): incluyen unidades más nuevas o con amenities, por lo que no son "
        "comparables directos. La brecha indica margen de ajuste en renovaciones, sujeto al estado de cada unidad.").font = NOTE

# ---------------------------------------------------------------- Plusvalia
PV = sheet("Plusvalia", "Plusvalía y retorno total a 5 años — tres escenarios", [52, 18, 18, 18])
legend(PV, 2)
header(PV, 4, ["Concepto", "Conservador", "Base", "Optimista"])
AP = [PR["a_cons"], PR["a_base"], PR["a_opt"]]
NETCOL = ["C", "E", "G"]
for k, col in enumerate("BCD"):
    put(PV, f"{col}5", f"={AP[k]}", LINK, "0.0%")
    put(PV, f"{col}6", f"={PR['precio']}*(1+{col}5)^5", CALC, USD)
    put(PV, f"{col}7", f"={col}6-{PR['precio']}", CALC, USD)
    put(PV, f"{col}8", f"=-{col}6*{PR['iva_venta']}", CALC, USD)
    put(PV, f"{col}9", f"=Proyeccion!{NETCOL[k]}10", LINK, USD)
    put(PV, f"{col}10", f"={col}7+{col}8+{col}9", BOLD, USD)
    put(PV, f"{col}11", f"={col}10/{PR['precio']}", CALC, "0.0%")
    # flujo para TIR
    put(PV, f"{col}14", f"=-{PR['precio']}", CALC, USD)
    for y in range(1, 6):
        extra = f"+{col}6+{col}8" if y == 5 else ""
        put(PV, f"{col}{14 + y}", f"=Proyeccion!{NETCOL[k]}{4 + y}{extra}", CALC, USD)
    put(PV, f"{col}20", f"=IRR({col}14:{col}19)", BOLD, "0.00%")
for i, lab in [(5, "Valorización anual supuesta"), (6, "Valor estimado al año 5"), (7, "Plusvalía estimada"), (8, "IVA de venta (1,5% efectivo)"),
               (9, "Renta neta acumulada 5 años"), (10, "Ganancia total estimada (plusvalía neta + renta)"), (11, "Ganancia total / precio"),
               (13, "Flujo para TIR (USD)"), (14, "Año 0 — compra"), (15, "Año 1"), (16, "Año 2"), (17, "Año 3"), (18, "Año 4"), (19, "Año 5 (incluye venta)"),
               (20, "TIR estimada a 5 años")]:
    put(PV, f"A{i}", lab, BOLD)
PV.cell(row=22, column=1, value="Escenarios hipotéticos: no constituyen garantía de plusvalía ni de rentabilidad. Sin gastos de adquisición ni comisión de venta. "
        "El valor de salida depende del mercado, del estado del edificio y de la ocupación al momento de vender.").font = NOTE

wb._sheets = [wb["Rentabilidad"], wb["Ingresos_Gastos"], wb["Superficies_Valor"], wb["Proyeccion"], wb["Plusvalia"], wb["Alquileres_Unidad"], wb["Conciliacion"], wb["Datos_Originales"], wb["Parametros"]]
for ws in wb.worksheets:
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(OUT)
