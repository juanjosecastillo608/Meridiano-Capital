"""Generador de cartas de oferta de Meridiano Capital (compra y alquiler).

Uso:
    python generar_carta_oferta.py <ficha.yaml> [--salida <carpeta>] [--sin-pdf]

Lee una ficha YAML (ver ejemplos/ y README.md), valida lo que se puede
validar automáticamente y genera en la carpeta de salida:

    02_Carta_Oferta_<codigo>_v<version>.docx
    02_Carta_Oferta_<codigo>_v<version>.pdf     (vía LibreOffice, si está)
    03_Control_Oferta_<codigo>_v<version>.md    (controles automáticos + observaciones)

Nunca sobrescribe una versión existente: si el archivo de esa versión ya existe,
se detiene y pide subir `version` en la ficha.

Reglas de contenido (prompt maestro de carta de oferta):
- No inventa datos: lo que falta se escribe como [PENDIENTE: ...] y se resalta.
- El estado se calcula: BORRADOR INCOMPLETO si queda algún [PENDIENTE] o falta un
  campo obligatorio; BORRADOR PARA REVISIÓN si está completo pero sin revisión
  jurídica marcada; LISTO PARA FIRMA solo con `revision_juridica_ok: true`.
Marca: lockup horizontal oficial (../logo_horizontal.png, D-034/D-036); paleta de
knowledge-base/brand/05-sistema-cromatico.md; Cambria/Calibri como alternativas
de sistema oficiales de Fraunces/Poppins (04-tipografia.md), sin negrita en el
wordmark serif (D-040).
"""
import argparse
import datetime as dt
import re
import shutil
import subprocess
import sys
from pathlib import Path

import yaml
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# ---------------------------------------------------------------- marca
PETROLEO = RGBColor(0x14, 0x31, 0x3A)
GRIS = RGBColor(0x6B, 0x61, 0x54)
TIERRA = RGBColor(0x8B, 0x33, 0x23)
DORADO_TXT = RGBColor(0xA8, 0x7D, 0x22)  # gold-d: dorado legible como texto sobre claro
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)
HEX_PETROLEO, HEX_DORADO, HEX_LINEA, HEX_FILA = "14313A", "C9982E", "DAD2C0", "F7F4EE"
SERIF, SANS = "Cambria", "Calibri"
PIE_DEFECTO = "Meridiano Capital · Asunción, Paraguay · meridianocapital.net"
# Lockup horizontal oficial vigente (D-034 + wordmark Fraunces D-036), rasterizado por ../rasterize_assets.js
LOGO = Path(__file__).resolve().parent.parent / "logo_horizontal.png"

PEND = re.compile(r"(\[PENDIENTE(?::[^\]]*)?\])")
MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto",
         "septiembre", "octubre", "noviembre", "diciembre"]


# ---------------------------------------------------------------- números
_U = ["cero", "uno", "dos", "tres", "cuatro", "cinco", "seis", "siete", "ocho", "nueve", "diez",
      "once", "doce", "trece", "catorce", "quince", "dieciséis", "diecisiete", "dieciocho",
      "diecinueve", "veinte", "veintiuno", "veintidós", "veintitrés", "veinticuatro", "veinticinco",
      "veintiséis", "veintisiete", "veintiocho", "veintinueve"]
_D = {3: "treinta", 4: "cuarenta", 5: "cincuenta", 6: "sesenta", 7: "setenta", 8: "ochenta", 9: "noventa"}
_C = {1: "ciento", 2: "doscientos", 3: "trescientos", 4: "cuatrocientos", 5: "quinientos",
      6: "seiscientos", 7: "setecientos", 8: "ochocientos", 9: "novecientos"}


def _menor_mil(n):
    if n == 100:
        return "cien"
    c, r = divmod(n, 100)
    partes = [_C[c]] if c else []
    if r:
        partes.append(_U[r] if r < 30 else _D[r // 10] + ("" if r % 10 == 0 else " y " + _U[r % 10]))
    return " ".join(partes)


def _apocope(t):
    return re.sub(r"veintiuno$", "veintiún", re.sub(r"(?<!veinti)uno$", "un", t))


def numero_a_letras(n: int) -> str:
    if n == 0:
        return "cero"
    millones, resto = divmod(n, 1_000_000)
    partes = []
    if millones:
        partes.append("un millón" if millones == 1 else _apocope(numero_a_letras(millones)) + " millones")
    miles, unidades = divmod(resto, 1000)
    if miles:
        partes.append("mil" if miles == 1 else _apocope(_menor_mil(miles)) + " mil")
    if unidades:
        partes.append(_menor_mil(unidades))
    return " ".join(partes)


def fmt_miles(n):
    return f"{n:,.0f}".replace(",", ".")


def fecha_larga(f: dt.date):
    return f"{f.day} de {MESES[f.month - 1]} de {f.year}"


# ---------------------------------------------------------------- docx helpers
class Carta:
    def __init__(self):
        self.doc = Document()
        sec = self.doc.sections[0]
        sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
        sec.top_margin, sec.bottom_margin = Cm(3.6), Cm(2.2)
        sec.left_margin = sec.right_margin = Cm(2.3)
        sec.header_distance, sec.footer_distance = Cm(1.0), Cm(1.0)
        st = self.doc.styles["Normal"]
        st.font.name, st.font.size = SANS, Pt(10.5)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), SANS)
        st.paragraph_format.space_after = Pt(5)
        st.paragraph_format.line_spacing = 1.12

    # -- texto con [PENDIENTE] resaltado
    @staticmethod
    def runs(p, texto, bold=False, color=None, size=None, font=None, italic=False, spacing=None):
        for trozo in PEND.split(str(texto)):
            if not trozo:
                continue
            r = p.add_run(trozo)
            r.bold, r.italic = bold, italic
            if font:
                r.font.name = font
            if size:
                r.font.size = Pt(size)
            if color is not None:
                r.font.color.rgb = color
            if spacing:
                rPr = r._r.get_or_add_rPr()
                sp = OxmlElement("w:spacing")
                sp.set(qn("w:val"), str(spacing))
                rPr.append(sp)
            if PEND.fullmatch(trozo):
                r.font.highlight_color = WD_COLOR_INDEX.YELLOW
                r.bold = False
        return p

    def p(self, texto="", align=None, after=None, before=None, keep=False, indent=None, **kw):
        par = self.doc.add_paragraph()
        self.runs(par, texto, **kw)
        if align is not None:
            par.alignment = align
        if after is not None:
            par.paragraph_format.space_after = Pt(after)
        if before is not None:
            par.paragraph_format.space_before = Pt(before)
        if keep:
            par.paragraph_format.keep_with_next = True
        if indent is not None:
            par.paragraph_format.first_line_indent = Cm(indent)
        return par

    @staticmethod
    def borde(par, lado="bottom", color=HEX_DORADO, sz=8):
        pPr = par._p.get_or_add_pPr()
        b = OxmlElement("w:pBdr")
        e = OxmlElement(f"w:{lado}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:space"), "4")
        e.set(qn("w:color"), color)
        b.append(e)
        pPr.append(b)

    @staticmethod
    def sombra(celda, fill):
        tcPr = celda._tc.get_or_add_tcPr()
        s = OxmlElement("w:shd")
        s.set(qn("w:val"), "clear")
        s.set(qn("w:color"), "auto")
        s.set(qn("w:fill"), fill)
        tcPr.append(s)

    @staticmethod
    def bordes_tabla(t, color=HEX_LINEA, ninguno=False):
        tblPr = t._tbl.tblPr
        b = OxmlElement("w:tblBorders")
        for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement(f"w:{lado}")
            e.set(qn("w:val"), "nil" if ninguno else "single")
            if not ninguno:
                e.set(qn("w:sz"), "4")
                e.set(qn("w:color"), color)
            b.append(e)
        tblPr.append(b)

    @staticmethod
    def fijar_anchos(t, anchos_cm):
        """Layout fijo + tblGrid explícito: LibreOffice ignora los anchos de celda si no."""
        t.autofit = False
        tblPr = t._tbl.tblPr
        lay = OxmlElement("w:tblLayout")
        lay.set(qn("w:type"), "fixed")
        tblPr.append(lay)
        for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), anchos_cm):
            gc.set(qn("w:w"), str(int(Cm(w).twips)))
        for fila in t.rows:
            for c, w in zip(fila.cells, anchos_cm):
                c.width = Cm(w)

    @staticmethod
    def no_cortar(fila):
        trPr = fila._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))

    def tabla_etiquetas(self, filas, anchos=(5.0, 11.4), oscura=False, size=9.5):
        """Tabla etiqueta | valor. oscura=True: etiquetas en petróleo (estilo bloque PRESENTA)."""
        t = self.doc.add_table(rows=len(filas), cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        self.bordes_tabla(t)
        self.fijar_anchos(t, anchos)
        for i, (et, val) in enumerate(filas):
            self.no_cortar(t.rows[i])
            c0, c1 = t.cell(i, 0), t.cell(i, 1)
            c0.width, c1.width = Cm(anchos[0]), Cm(anchos[1])
            for c in (c0, c1):
                c.paragraphs[0].paragraph_format.space_after = Pt(2)
                c.paragraphs[0].paragraph_format.space_before = Pt(2)
                c.paragraphs[0].paragraph_format.keep_with_next = i < len(filas) - 1
            self.runs(c0.paragraphs[0], et.upper(), bold=True, size=size - 0.5,
                      color=BLANCO if oscura else PETROLEO)
            self.runs(c1.paragraphs[0], val, size=size, bold=oscura)
            self.sombra(c0, HEX_PETROLEO if oscura else HEX_FILA)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)
        return t

    def tabla_grilla(self, filas, anchos, total=False, size=9):
        t = self.doc.add_table(rows=len(filas), cols=len(filas[0]))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        self.bordes_tabla(t)
        self.fijar_anchos(t, anchos)
        for i, fila in enumerate(filas):
            self.no_cortar(t.rows[i])
            cab, tot = i == 0, total and i == len(filas) - 1
            for j, val in enumerate(fila):
                c = t.cell(i, j)
                c.width = Cm(anchos[j])
                c.paragraphs[0].paragraph_format.space_after = Pt(1)
                c.paragraphs[0].paragraph_format.keep_with_next = i < len(filas) - 1
                self.runs(c.paragraphs[0], val, size=size, bold=cab or tot,
                          color=BLANCO if cab else (PETROLEO if tot else None))
                if cab:
                    self.sombra(c, HEX_PETROLEO)
                elif i % 2 == 0:
                    self.sombra(c, HEX_FILA)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)
        return t

    def membrete(self, logo: Path, pie: str):
        sec = self.doc.sections[0]
        hdr = sec.header.paragraphs[0]
        hdr.add_run().add_picture(str(logo), height=Cm(2.0))
        hdr.paragraph_format.space_after = Pt(2)
        self.borde(hdr, color=HEX_DORADO, sz=6)
        f = sec.footer.paragraphs[0]
        self.borde(f, lado="top", color=HEX_LINEA, sz=6)
        self.runs(f, pie, size=8, color=GRIS)

    def firmas(self, firmas):
        t = self.doc.add_table(rows=1, cols=len(firmas))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        self.bordes_tabla(t, ninguno=True)
        ancho = 16.4 / len(firmas)
        self.fijar_anchos(t, [ancho] * len(firmas))
        for j, f in enumerate(firmas):
            c = t.cell(0, j)
            c.width = Cm(ancho)
            par = c.paragraphs[0]
            par.paragraph_format.space_before = Pt(42)
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            self.runs(par, "______________________________", color=GRIS)
            for k, (clave, estilo) in enumerate((("nombre", dict(bold=True, color=PETROLEO)),
                                                 ("documento", dict(color=GRIS, size=9)),
                                                 ("cargo", dict(color=GRIS, size=9)),
                                                 ("entidad", dict(bold=True, color=DORADO_TXT, size=9)))):
                if f.get(clave):
                    q = c.add_paragraph()
                    q.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    q.paragraph_format.space_after = Pt(0)
                    self.runs(q, f[clave], **estilo)
        self.no_cortar(t.rows[0])


# ---------------------------------------------------------------- validación
OBLIGATORIOS = {
    "comun": ["codigo", "version", "tipo", "rol_meridiano", "fecha", "destinatario", "presenta",
              "proponente", "referencia", "vigencia", "efecto_aceptacion", "intervencion", "firmas"],
    "compra": ["precio", "pagos"],
    "alquiler": ["canon"],
}


def validar(f):
    obs, bloqueos = [], []
    for k in OBLIGATORIOS["comun"] + OBLIGATORIOS.get(f.get("tipo"), []):
        if f.get(k) in (None, "", []):
            bloqueos.append(f"Falta el campo obligatorio `{k}`.")
    if f.get("tipo") not in ("compra", "alquiler"):
        bloqueos.append("`tipo` debe ser `compra` o `alquiler`.")
    if f.get("rol_meridiano") not in ("CAPTADOR", "COLOCADOR", "AMBOS"):
        bloqueos.append("`rol_meridiano` debe ser CAPTADOR, COLOCADOR o AMBOS.")

    def textos(x):
        if isinstance(x, dict):
            for v in x.values():
                yield from textos(v)
        elif isinstance(x, list):
            for v in x:
                yield from textos(v)
        elif isinstance(x, str):
            yield x

    sin_obs = {k: v for k, v in f.items() if k != "observaciones"}
    texto = "\n".join(textos(sin_obs))
    pendientes = sorted(set(PEND.findall(texto)))
    if re.search(r"en representaci[oó]n de", texto, re.I) and not f.get("mandato_documentado"):
        obs.append("⚠️ El texto dice \"en representación de\" sin `mandato_documentado: true`. "
                   "Reemplazar por captador/colocador o adjuntar la autorización.")

    if f.get("tipo") == "compra" and f.get("precio"):
        pr = f["precio"]
        imp = pr.get("ofertado_importe")
        if isinstance(imp, int):
            esperado = numero_a_letras(imp)
            letras = (pr.get("ofertado_letras") or "").lower()
            if esperado in letras:
                obs.append(f"✅ Cifras vs. letras del precio ofertado: {fmt_miles(imp)} = \"{esperado}\".")
            else:
                bloqueos.append(f"Cifras vs. letras NO coinciden: {fmt_miles(imp)} debería leerse \"{esperado}\", "
                                f"la ficha dice \"{letras}\".")
            numericos = [p["importe"] for p in f.get("pagos", []) if isinstance(p.get("importe"), int)]
            variables = [p for p in f.get("pagos", []) if not isinstance(p.get("importe"), int)]
            if not variables:
                if sum(numericos) == imp:
                    obs.append(f"✅ Suma del cronograma = precio ofertado ({fmt_miles(imp)}).")
                else:
                    bloqueos.append(f"Suma del cronograma ({fmt_miles(sum(numericos))}) ≠ precio ofertado ({fmt_miles(imp)}).")
            else:
                ok = all(p.get("complementa_al_total") for p in variables) and len(variables) <= 2
                obs.append(("✅" if ok else "⚠️") + " El cronograma tiene importes a convenir "
                           f"({', '.join(p['concepto'] for p in variables)}); el total queda fijado en "
                           f"{fmt_miles(imp)} y los parciales deben complementarse hasta ese total.")
        com = f.get("comision")
        if com and isinstance(imp, int):
            monto = imp * com["porcentaje"] / 100
            if com.get("iva_incluido"):
                neto = monto / 1.1
                obs.append(f"🧮 Comisión {com['porcentaje']} % IVA incluido sobre {fmt_miles(imp)}: "
                           f"{pr.get('moneda', '')} {fmt_miles(monto)} (neto {fmt_miles(neto)} + IVA {fmt_miles(monto - neto)}), "
                           f"a cargo de {com.get('a_cargo_de', '[PENDIENTE]')}.")
            else:
                obs.append(f"🧮 Comisión {com['porcentaje']} % + IVA sobre {fmt_miles(imp)}: "
                           f"{fmt_miles(monto)} + IVA {fmt_miles(monto * 0.1)} = {fmt_miles(monto * 1.1)}.")

    vig = f.get("vigencia") or {}
    if vig.get("dias") and vig.get("fecha_presentacion"):
        fp = vig["fecha_presentacion"]
        if vig.get("tipo", "corridos") == "corridos":
            vence = fp + dt.timedelta(days=vig["dias"])
            obs.append(f"📅 Vigencia: {vig['dias']} días corridos desde {fp:%d-%m-%Y} → vence el {vence:%d-%m-%Y} "
                       f"({vig.get('hora') or 'fin del día'}, {vig.get('zona', 'America/Asuncion')}).")
        else:
            obs.append("⚠️ Vigencia en días hábiles: calcular a mano considerando feriados de Paraguay.")
    elif vig.get("dias"):
        obs.append(f"📅 Vigencia: {vig['dias']} días {vig.get('tipo', 'corridos')} desde {vig.get('desde', 'la presentación')}; "
                   "la fecha exacta queda determinada por la constancia de recepción.")

    if bloqueos or pendientes:
        estado = "BORRADOR INCOMPLETO"
    elif not f.get("revision_juridica_ok"):
        estado = "BORRADOR PARA REVISIÓN"
    else:
        estado = "LISTO PARA FIRMA"
    return estado, obs, bloqueos, pendientes


# ---------------------------------------------------------------- render
def vigencia_texto(v):
    if v.get("texto"):
        return v["texto"]
    base = (f"La presente oferta tendrá una vigencia de {v['dias']} ({numero_a_letras(v['dias'])}) días "
            f"{v.get('tipo', 'corridos')} contados desde {v.get('desde', 'la fecha de su presentación')}")
    if v.get("hora"):
        base += f", hasta las {v['hora']} (hora de Asunción)"
    base += ". Vencido dicho plazo sin aceptación expresa por escrito, la oferta quedará sin efecto."
    if v.get("canal"):
        base += " " + v["canal"]
    return base


def render(f, salida: Path, estado):
    c = Carta()
    c.membrete(LOGO, f.get("pie", PIE_DEFECTO))

    es_compra = f["tipo"] == "compra"
    if estado != "LISTO PARA FIRMA":
        c.p(f"{estado} — documento de trabajo, no apto para firma"
            + (" hasta resolver los puntos resaltados" if estado == "BORRADOR INCOMPLETO" else " hasta completar la revisión jurídica"),
            color=TIERRA, size=8.5, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)

    c.p(f.get("titulo", "CARTA DE OFERTA DE COMPRA" if es_compra else "CARTA DE OFERTA DE ALQUILER"),
        font=SERIF, size=16, color=PETROLEO, align=WD_ALIGN_PARAGRAPH.CENTER, after=0, spacing=20)
    sub = c.p(f.get("subtitulo", ""), italic=True, color=PETROLEO, size=10.5,
              align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
    c.borde(sub, color=HEX_PETROLEO, sz=6)

    fecha = f["fecha"]
    c.p(f"{f.get('ciudad', 'Asunción')}, {fecha_larga(fecha)}", color=GRIS, align=WD_ALIGN_PARAGRAPH.RIGHT, after=0)
    c.p(f"Ref. interna: {f['codigo']} · v{f['version']}", color=GRIS, size=8, align=WD_ALIGN_PARAGRAPH.RIGHT, after=8)

    prop = f["proponente"]
    prop_txt = " · ".join(x for x in (prop.get("nombre"), prop.get("documento"),
                                       f"Representante: {prop['representante']}" if prop.get("representante") else None) if x)
    bloque = [("Presenta", f["presenta"]), (prop.get("etiqueta", "Proponente"), prop_txt)]
    bloque += [tuple(x) for x in f.get("bloque_partes_extra", [])]
    c.tabla_etiquetas(bloque, anchos=(5.2, 11.2), oscura=True)

    for i, linea in enumerate(f["destinatario"]):
        c.p(linea, bold=i not in (0,) and linea.lower() != "presente", after=0)
    c.p("", after=4)
    c.p(f"Ref.: {f['referencia']}", bold=True, color=PETROLEO, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)

    for par in f.get("parrafos_intro", []):
        c.p(par, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=1.0)

    if f.get("inmueble"):
        c.p(f.get("titulo_inmueble", "Inmueble"), font=SERIF, size=12, color=PETROLEO, keep=True, before=4, after=3)
        c.tabla_etiquetas([tuple(x) for x in f["inmueble"]])

    if es_compra:
        pr = f["precio"]
        c.p("Precio", font=SERIF, size=12, color=PETROLEO, keep=True, before=4, after=3)
        filas = []
        if pr.get("solicitado"):
            filas.append(("Precio de venta publicado (referencia)", pr["solicitado"]))
        imp = pr["ofertado_importe"]
        filas.append(("Precio ofertado", f"{pr.get('moneda', 'Gs')} {fmt_miles(imp) if isinstance(imp, int) else imp} "
                                         f"({pr.get('ofertado_letras', '[PENDIENTE: en letras]')})"))
        c.tabla_etiquetas(filas)
        if pr.get("nota"):
            c.p(pr["nota"], align=WD_ALIGN_PARAGRAPH.JUSTIFY)

        c.p("Forma de pago", font=SERIF, size=12, color=PETROLEO, keep=True, before=4, after=3)
        filas = [["Concepto", "Importe", "Moneda", "Fecha o hito", "Condición"]]
        for pg in f["pagos"]:
            v = pg.get("importe")
            filas.append([pg["concepto"], fmt_miles(v) if isinstance(v, int) else (v or "[PENDIENTE]"),
                          pg.get("moneda", pr.get("moneda", "Gs")), pg.get("hito", ""), pg.get("condicion", "")])
        filas.append(["Total", fmt_miles(imp) if isinstance(imp, int) else str(imp), pr.get("moneda", "Gs"), "", ""])
        c.tabla_grilla(filas, [3.9, 3.2, 1.7, 3.8, 3.8], total=True)
        for nota in f.get("notas_pago", []):
            c.p(nota, size=9.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    else:
        c.p(f.get("titulo_canon", "Condiciones económicas"), font=SERIF, size=12, color=PETROLEO, keep=True, before=4, after=3)
        c.tabla_etiquetas([tuple(x) for x in f["canon"]])

    if f.get("condiciones"):
        c.p(f.get("titulo_condiciones", "Condiciones de la oferta"), font=SERIF, size=12, color=PETROLEO,
            keep=True, before=4, after=3)
        for i, cond in enumerate(f["condiciones"], 1):
            par = c.p("", after=4, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
            par.paragraph_format.left_indent = Cm(0.6)
            par.paragraph_format.first_line_indent = Cm(-0.6)
            c.runs(par, f"{i}.  ", color=PETROLEO)
            if cond.get("titulo"):
                c.runs(par, cond["titulo"] + ". ", bold=True)
            c.runs(par, cond["texto"])

    c.p("Vigencia, respuesta y alcance", font=SERIF, size=12, color=PETROLEO, keep=True, before=4, after=3)
    c.p(vigencia_texto(f["vigencia"]), align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    c.p(f["efecto_aceptacion"], align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    c.p(f["intervencion"], align=WD_ALIGN_PARAGRAPH.JUSTIFY, color=GRIS, keep=True)
    c.p(f.get("cierre", "Sin otro particular y aguardando una respuesta favorable, saludamos atentamente."), keep=True)
    c.firmas(f["firmas"])

    resp = f.get("respuesta") or {}
    if f.get("recepcion", True) or resp.get("habilitada", True):
        t = c.p("Constancia de recepción", font=SERIF, size=13, color=PETROLEO, after=2)
        t.paragraph_format.page_break_before = True
        c.borde(t, color=HEX_LINEA, sz=6)
        c.p(f.get("texto_recepcion", "Recibí copia de esta oferta. Esta constancia acredita únicamente la recepción "
                                     "y no implica aceptación. La fecha consignada aquí es la fecha de presentación "
                                     "a partir de la cual corre la vigencia."), size=10)
        c.tabla_etiquetas([("Fecha y hora de recepción", "____ / ____ / ______  ·  ____ : ____ hs (hora de Asunción)"),
                           ("Firma", ""), ("Aclaración y documento", "")], anchos=(5.2, 11.2))
    if resp.get("habilitada", True):
        t = c.p(f"Respuesta de la {resp.get('parte', 'parte destinataria')}", font=SERIF, size=13,
                color=PETROLEO, before=8, after=2)
        c.borde(t, color=HEX_LINEA, sz=6)
        c.p("Marcar una sola opción:", size=10)
        for op in ["☐  Acepto íntegramente la oferta en los términos de esta carta.",
                   "☐  Rechazo la oferta.",
                   "☐  Formulo la siguiente contraoferta (una contraoferta no implica aceptación de los términos originales):"]:
            c.p(op, after=3)
        caja = c.doc.add_table(rows=1, cols=1)
        c.bordes_tabla(caja)
        c.fijar_anchos(caja, [16.4])
        caja.rows[0].height = Cm(2.6)
        c.doc.add_paragraph().paragraph_format.space_after = Pt(0)
        if resp.get("texto_aceptacion"):
            c.p(resp["texto_aceptacion"], size=10, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        c.tabla_etiquetas([(x, "") if isinstance(x, str) else tuple(x) for x in
                           resp.get("campos", ["Firma", "Aclaración y documento",
                                               ["Fecha y hora", "____ / ____ / ______  ·  ____ : ____ hs (hora de Asunción)"]])],
                          anchos=(5.2, 11.2))

    base = f"02_Carta_Oferta_{f['codigo']}_v{f['version']}"
    docx = salida / f"{base}.docx"
    c.doc.save(docx)
    return docx


def control_md(f, estado, obs, bloqueos, pendientes, docx, pdf):
    L = [f"# Control automático — {f['codigo']} v{f['version']}", "",
         f"**Estado calculado: {estado}**", "",
         f"Generado por `production/generadores/cartas_oferta/generar_carta_oferta.py` el {dt.date.today():%d-%m-%Y}. "
         "Documento interno — no se envía a las partes.", "",
         f"- Tipo: {f.get('tipo')} · Rol de Meridiano: {f.get('rol_meridiano')}",
         f"- Carta: `{docx.name}`" + (f" · `{pdf.name}`" if pdf else " · PDF no generado"), "",
         "## Bloqueos", ""]
    L += [f"- ❌ {b}" for b in bloqueos] or ["- Ninguno."]
    L += ["", "## Controles automáticos", ""] + [f"- {o}" for o in obs]
    L += ["", f"## Marcadores [PENDIENTE] en la carta ({len(pendientes)})", ""]
    L += [f"- {p}" for p in pendientes] or ["- Ninguno."]
    if f.get("observaciones"):
        L += ["", "## Observaciones del analista (de la ficha)", ""] + [f"- {o}" for o in f["observaciones"]]
    L += ["", "## Recordatorio", "",
          "- Revisar visualmente el PDF (tablas, firmas, saltos) antes de enviarlo.",
          "- Correr el barrido de PII de `CLAUDE.md` antes de commitear dentro de `contracts/`.",
          "- LISTO PARA FIRMA no significa firmado ni aceptado."]
    return "\n".join(L) + "\n"


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("ficha")
    ap.add_argument("--salida")
    ap.add_argument("--sin-pdf", action="store_true")
    a = ap.parse_args()

    ficha = Path(a.ficha)
    f = yaml.safe_load(ficha.read_text(encoding="utf-8"))
    salida = Path(a.salida) if a.salida else ficha.parent
    salida.mkdir(parents=True, exist_ok=True)
    base = f"02_Carta_Oferta_{f.get('codigo')}_v{f.get('version')}"
    if (salida / f"{base}.docx").exists():
        sys.exit(f"Ya existe {base}.docx — subir `version` en la ficha para no sobrescribir una versión previa.")

    estado, obs, bloqueos, pendientes = validar(f)
    if any("obligatorio" in b or "debe ser" in b for b in bloqueos):
        sys.exit("No se puede generar:\n- " + "\n- ".join(bloqueos))
    docx = render(f, salida, estado)

    pdf = None
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice and not a.sin_pdf:
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", str(docx), "--outdir", str(salida)],
                       check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        pdf = docx.with_suffix(".pdf")
    ctrl = salida / f"03_Control_Oferta_{f['codigo']}_v{f['version']}.md"
    ctrl.write_text(control_md(f, estado, obs, bloqueos, pendientes, docx, pdf), encoding="utf-8")
    print(f"{estado}\n{docx}\n{pdf or '(sin PDF)'}\n{ctrl}")


if __name__ == "__main__":
    main()
